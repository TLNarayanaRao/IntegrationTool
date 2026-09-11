from __future__ import annotations
import ast, asyncio, base64, csv, ftplib, gzip, importlib.util, io, json, os, re, shlex, shutil, sqlite3, sys, tempfile, traceback, uuid
from datetime import datetime, timezone, timedelta
from pathlib import Path
import httpx
from time import perf_counter
from .models import Activity, GroupDefinition, ProcessDefinition, Project, RunResult
from .mapper import apply_function, execute as execute_mapping
from .dataweave import DataWeaveError, execute as execute_dataweave
from .sap import sap_adapter
from .snowflake import snowflake_adapter
from .jdbc import jdbc_adapter
from .amqp import amqp_adapter
from .java_bridge import JavaBridgeError, execute_jms
from .google_pubsub import client_configuration as pubsub_client_configuration, create_client as create_pubsub_client
from .time_utils import log_timestamp

class RuntimeErrorWithLogs(Exception): pass
_NO_EVENT_OUTPUT = object()
class FabricFault(Exception):
    def __init__(self, message: str, *, fault_type='UserDefinedException', code='', details=None, cause=None):
        super().__init__(message); self.fault_type = fault_type; self.code = code; self.details = details or {}; self.cause = cause

class WorkflowRuntime:
    def __init__(self):
        self.messages: dict[str, list[dict]] = {}
        self.acknowledgements: dict[str, dict] = {}
        self.shared_variables: dict[str, Any] = {}
        self.group_locks: dict[str, asyncio.Lock] = {}

    def register_acknowledgement(self, technology: str, message_id: str, callback=None) -> str:
        ack_id = f'{technology}:{message_id}:{uuid.uuid4()}'
        self.acknowledgements[ack_id] = {'technology': technology, 'messageId': message_id, 'callback': callback, 'created': log_timestamp()}
        return ack_id

    async def confirm_messages(self, handles) -> dict:
        handles = handles if isinstance(handles, list) else [handles]
        confirmed, technologies = [], []
        for handle in [item for item in handles if item]:
            pending = self.acknowledgements.pop(str(handle), None)
            if not pending: raise RuntimeError(f'Acknowledgement handle {handle!r} was not found or was already confirmed')
            callback = pending.get('callback')
            if callback:
                result = callback()
                if asyncio.iscoroutine(result): await result
            confirmed.append(str(handle)); technologies.append(pending['technology'])
        return {'confirmed': True, 'count': len(confirmed), 'ackIds': confirmed, 'technologies': sorted(set(technologies))}

    def compile_groups(self, process: ProcessDefinition) -> dict[str, dict]:
        """Compile persisted group containers into executable graph boundaries."""
        groups = {group.id: group for group in process.groups}
        children = {group_id: [] for group_id in groups}
        for group in process.groups:
            if group.parent_group_id: children[group.parent_group_id].append(group.id)

        def descendants(group_id: str) -> set[str]:
            result = set(groups[group_id].member_activity_ids)
            for child_id in children[group_id]: result.update(descendants(child_id))
            return result

        plans: dict[str, dict] = {}
        for group in process.groups:
            config = group.config or {}
            if group.type in ('if', 'while') and not str(config.get('condition') or '').strip():
                raise FabricFault(f'Group {group.name} requires a boolean condition', fault_type='GROUP_VALIDATION')
            if group.type in ('iterate', 'for_each') and config.get('source') in (None, '') and config.get('collection') in (None, '') and not (group.type == 'for_each' and config.get('start') is not None and config.get('end') is not None):
                raise FabricFault(f'Group {group.name} requires a collection expression', fault_type='GROUP_VALIDATION')
            if group.type == 'repeat' and not str(config.get('condition') or '').strip() and config.get('count') is None and config.get('iterations') is None:
                raise FabricFault(f'Group {group.name} requires a Repeat Until True condition', fault_type='GROUP_VALIDATION')
            if group.type == 'repeat_on_error' and not str(config.get('stopCondition') or '').strip():
                raise FabricFault(f'Group {group.name} requires a Repeat-on-Error stop condition', fault_type='GROUP_VALIDATION')
            members = descendants(group.id)
            if not members: raise FabricFault(f'Group {group.name} is empty', fault_type='GROUP_VALIDATION')
            internal_incoming = {edge.target for edge in process.transitions if edge.source in members and edge.target in members}
            external_entries = {edge.target for edge in process.transitions if edge.source not in members and edge.target in members}
            entries = external_entries or (members - internal_incoming)
            exit_edges = [edge for edge in process.transitions if edge.source in members and edge.target not in members]
            for source_id in members:
                fanout = [edge for edge in process.transitions if edge.source == source_id and edge.target in members and edge.type == 'success']
                if len(fanout) > 1 and group.type != 'pick_first':
                    raise FabricFault(f'Group {group.name} contains a parallel fan-out; use separate top-level branches until grouped branch joining is qualified', fault_type='GROUP_UNSUPPORTED')
            exit_sources = {edge.source for edge in exit_edges}
            terminal = {activity.id for activity in process.activities if activity.id in members and not any(edge.source == activity.id and edge.target in members for edge in process.transitions)}
            exits = exit_sources or terminal
            if len(entries) != 1 and group.type != 'pick_first':
                raise FabricFault(f'Group {group.name} must have exactly one entry activity; found {len(entries)}', fault_type='GROUP_VALIDATION')
            if len(exits) != 1 and group.type != 'pick_first':
                raise FabricFault(f'Group {group.name} must have exactly one exit activity; found {len(exits)}', fault_type='GROUP_VALIDATION')
            plans[group.id] = {'group': group, 'members': members, 'entries': entries, 'entry': next(iter(entries)) if entries else None, 'exits': exits, 'exit': next(iter(exits)) if exits else None, 'exitEdges': exit_edges}
        return plans

    @staticmethod
    def _group_ancestors(group: GroupDefinition, plans: dict[str, dict]) -> list[str]:
        result, current = [], group
        while current:
            result.append(current.id)
            current = plans.get(current.parent_group_id, {}).get('group') if current.parent_group_id else None
        return list(reversed(result))

    @staticmethod
    def _set_group_variable(state: dict, ctx: dict, name: str, value) -> None:
        if not name: return
        previous = state.setdefault('previousVariables', {})
        if name not in previous: previous[name] = (name in ctx['vars'], ctx['vars'].get(name))
        ctx['vars'][name] = value

    @staticmethod
    def _restore_group_variables(state: dict, ctx: dict) -> None:
        for name, (existed, value) in state.get('previousVariables', {}).items():
            if existed: ctx['vars'][name] = value
            else: ctx['vars'].pop(name, None)

    @staticmethod
    def _publish_group_context(state: dict, group: GroupDefinition, ctx: dict, current_element=None) -> None:
        details = {'id': group.id, 'name': group.name, 'type': group.type, 'iteration': state.get('iteration', 0) + 1}
        if current_element is not None: details['currentElement'] = current_element
        ctx['context']['group'] = details

    async def _begin_group(self, plan: dict, ctx: dict) -> tuple[dict, bool]:
        # Keep condition expressions intact for the condition evaluator. Exact
        # value fields are resolved individually below.
        group, cfg = plan['group'], dict(plan['group'].config)
        state = {'id': group.id, 'iteration': 0, 'config': cfg}
        kind = group.type
        should_run = True
        if kind in ('if', 'while'):
            expression = str(cfg.get('condition') or '')
            if not expression: raise FabricFault(f'{group.name} requires a condition', fault_type='GROUP_VALIDATION')
            should_run = self.condition(expression, ctx)
            if kind == 'while':
                self._set_group_variable(state, ctx, str(cfg.get('indexVariable') or 'index'), 1)
                self._set_group_variable(state, ctx, 'currentIndex', 1)
        elif kind in ('for_each', 'iterate'):
            source_value = cfg.get('collection', cfg.get('source'))
            if kind == 'for_each' and source_value in (None, ''):
                start, end, increment = int(self.resolve(cfg.get('start', 1), ctx)), int(self.resolve(cfg.get('end', 1), ctx)), int(self.resolve(cfg.get('increment', 1), ctx) or 1)
                if increment == 0: raise FabricFault(f'{group.name} increment cannot be zero', fault_type='GROUP_VALIDATION')
                items = list(range(start, end + (1 if increment > 0 else -1), increment))
            else:
                source = self.resolve(source_value or [], ctx)
                items = list(source.values()) if isinstance(source, dict) else list(source or [])
            state['items'] = items; should_run = bool(items)
            if should_run:
                current_name = str(cfg.get('currentElementName') or cfg.get('itemVariable') or 'currentElement')
                self._set_group_variable(state, ctx, current_name, items[0])
                self._set_group_variable(state, ctx, 'currentElement', items[0])
                self._set_group_variable(state, ctx, str(cfg.get('indexVariable') or 'index'), 1)
                self._set_group_variable(state, ctx, 'currentIndex', 1)
                state['currentElementName'] = current_name
            if self.as_bool(cfg.get('accumulateOutput', False)):
                state['accumulatedOutput'] = []
                state['accumulatorVariable'] = str(cfg.get('accumulatorVariable') or f'{group.id}Results')
                ctx['vars'][state['accumulatorVariable']] = state['accumulatedOutput']
        elif kind == 'repeat':
            state['count'] = max(0, int(self.resolve(cfg.get('count', cfg.get('iterations', 1)), ctx) or 0)); should_run = True if cfg.get('condition') else state['count'] > 0
            self._set_group_variable(state, ctx, str(cfg.get('indexVariable') or 'index'), 1)
            self._set_group_variable(state, ctx, 'currentIndex', 1)
        elif kind == 'repeat_on_error':
            state['retriesRemaining'] = max(0, int(self.resolve(cfg.get('retryCount', cfg.get('retries', 3)), ctx) or 0))
            self._set_group_variable(state, ctx, str(cfg.get('indexVariable') or 'index'), 1)
            self._set_group_variable(state, ctx, 'currentIndex', 1)
        elif kind == 'critical_section':
            lock_name = str(self.resolve(cfg.get('lockName'), ctx) or f'{ctx.get("_process").id}:{group.id}')
            lock = self.group_locks.setdefault(lock_name, asyncio.Lock())
            await lock.acquire(); state['lock'] = lock; state['lockName'] = lock_name
        elif kind == 'transaction_jdbc':
            resource_id = str(self.resolve(cfg.get('resourceId'), ctx) or '')
            jdbc_members = [item for item in ctx.get('_process').activities if item.id in plan['members'] and item.type == 'jdbc']
            inferred = {str(self.resolve(item.config.get('resourceId'), ctx) or '') for item in jdbc_members} - {''}
            if not resource_id and len(inferred) == 1: resource_id = next(iter(inferred))
            if not resource_id or len(inferred - {resource_id}) > 0:
                raise FabricFault(f'{group.name} must use one JDBC shared connection', fault_type='JDBCTransactionException')
            resource = ctx['resources'].get(resource_id)
            if not resource or resource.type != 'jdbc': raise FabricFault(f'{group.name} requires a valid JDBC connection', fault_type='JDBCTransactionException')
            connection_config = self.resolve(resource.config, ctx)
            state['resourceId'] = resource_id
            state['connection'] = await asyncio.to_thread(jdbc_adapter.connect, connection_config)
            ctx.setdefault('jdbcTransactions', {})[resource_id] = state['connection']
        self._publish_group_context(state, group, ctx, state.get('items', [None])[0] if state.get('items') else None)
        self.log(ctx['logs'], 'DEBUG', f'Group entered: {group.name}', kind='group', groupId=group.id, groupType=group.type, iteration=0)
        return state, should_run

    async def _finish_group(self, state: dict, plan: dict, ctx: dict, success: bool) -> None:
        if plan['group'].type == 'transaction_jdbc':
            connection = state.get('connection')
            if connection:
                try: await asyncio.to_thread(connection.commit if success else connection.rollback)
                finally: await asyncio.to_thread(connection.close)
            ctx.setdefault('jdbcTransactions', {}).pop(state.get('resourceId'), None)
        if plan['group'].type == 'critical_section' and state.get('lock') and state['lock'].locked(): state['lock'].release()
        self._restore_group_variables(state, ctx)
        self.log(ctx['logs'], 'DEBUG', f'Group {"completed" if success else "failed"}: {plan["group"].name}', kind='group', groupId=plan['group'].id, groupType=plan['group'].type, iteration=state.get('iteration', 0))

    async def enter_group_boundaries(self, activity_id: str, ctx: dict, plans: dict[str, dict]) -> str | None:
        active = {state['id'] for state in ctx.setdefault('groupStack', [])}
        candidates = [plan for plan in plans.values() if activity_id in plan.get('entries', {plan.get('entry')}) and plan['group'].id not in active]
        candidates.sort(key=lambda plan: len(self._group_ancestors(plan['group'], plans)))
        for plan in candidates:
            if plan['group'].parent_group_id and plan['group'].parent_group_id not in {state['id'] for state in ctx['groupStack']}: continue
            state, should_run = await self._begin_group(plan, ctx)
            ctx['groupStack'].append(state)
            if not should_run:
                await self._finish_group(state, plan, ctx, True); ctx['groupStack'].pop()
                edges = plan['exitEdges']
                eligible = self.select_group_transitions(self.eligible_success_transitions(edges, ctx), ctx, plans)
                return eligible[0].target if eligible else None
        return activity_id

    async def leave_group_boundaries(self, source_id: str, target_id: str | None, ctx: dict, plans: dict[str, dict], success: bool = True) -> str | None:
        while ctx.get('groupStack'):
            state = ctx['groupStack'][-1]; plan = plans[state['id']]; group = plan['group']
            if target_id in plan['members']: break
            if success and (source_id == plan['exit'] or source_id in plan.get('exits', {plan.get('exit')})):
                cfg = state['config']; repeat = False
                if group.type in ('for_each', 'iterate') and 'accumulatedOutput' in state:
                    state['accumulatedOutput'].append(ctx.get('last'))
                    ctx['vars'][state['accumulatorVariable']] = list(state['accumulatedOutput'])
                if group.type == 'while': repeat = self.condition(str(cfg.get('condition') or ''), ctx)
                elif group.type == 'repeat': repeat = (not self.condition(str(cfg.get('condition')), ctx)) if cfg.get('condition') else state['iteration'] + 1 < state['count']
                elif group.type in ('for_each', 'iterate'):
                    repeat = state['iteration'] + 1 < len(state['items'])
                if repeat:
                    state['iteration'] += 1
                    if group.type in ('for_each', 'iterate'):
                        current_element = state['items'][state['iteration']]
                        self._set_group_variable(state, ctx, state.get('currentElementName') or str(cfg.get('itemVariable') or 'currentElement'), current_element)
                        self._set_group_variable(state, ctx, 'currentElement', current_element)
                    if group.type in ('for_each', 'iterate', 'repeat', 'while'):
                        self._set_group_variable(state, ctx, str(cfg.get('indexVariable') or 'index'), state['iteration'] + 1)
                        self._set_group_variable(state, ctx, 'currentIndex', state['iteration'] + 1)
                    maximum = max(1, int(self.resolve(cfg.get('maxIterations', 10000), ctx) or 10000))
                    if state['iteration'] >= maximum: raise FabricFault(f'{group.name} exceeded maxIterations={maximum}', fault_type='GROUP_ITERATION_LIMIT')
                    self._publish_group_context(state, group, ctx, state['items'][state['iteration']] if group.type in ('for_each', 'iterate') else None)
                    self.log(ctx['logs'], 'DEBUG', f'Group iteration: {group.name} #{state["iteration"] + 1}', kind='group', groupId=group.id, groupType=group.type, iteration=state['iteration'])
                    return plan['entry']
            await self._finish_group(state, plan, ctx, success); ctx['groupStack'].pop()
        return target_id

    async def retry_failed_group(self, ctx: dict, plans: dict[str, dict]) -> str | None:
        stack = ctx.get('groupStack', [])
        retry_index = next((index for index in range(len(stack) - 1, -1, -1) if plans[stack[index]['id']]['group'].type == 'repeat_on_error' and stack[index].get('retriesRemaining', 0) > 0), None)
        if retry_index is None: return None
        while len(stack) - 1 > retry_index:
            child = stack.pop(); await self._finish_group(child, plans[child['id']], ctx, False)
        state = stack[retry_index]; state['retriesRemaining'] -= 1; state['iteration'] += 1
        if state['config'].get('stopCondition') and self.condition(str(state['config']['stopCondition']), ctx): return None
        self._set_group_variable(state, ctx, str(state['config'].get('indexVariable') or 'index'), state['iteration'] + 1)
        self._set_group_variable(state, ctx, 'currentIndex', state['iteration'] + 1)
        group = plans[state['id']]['group']
        self._publish_group_context(state, group, ctx)
        for activity_id in plans[state['id']]['members']:
            ctx.get('activities', {}).pop(activity_id, None)
            task_id = ctx.get('context', {}).get('taskId')
            if task_id: ctx.get('tasks', {}).get(task_id, {}).get('activities', {}).pop(activity_id, None)
        delay = float(self.resolve(state['config'].get('retryIntervalSeconds', state['config'].get('retryDelaySeconds', 0)), ctx) or 0)
        self.log(ctx['logs'], 'WARN', f'Group retry: {plans[state["id"]]["group"].name}; {state["retriesRemaining"]} retries remain', kind='group', groupId=state['id'], iteration=state['iteration'])
        if delay: await asyncio.sleep(delay)
        return plans[state['id']]['entry']

    async def run(self, process: ProcessDefinition, initial: dict, resources=None, properties=None, entry_activity_id=None, project: Project | None=None, execution_state: dict | None=None, event_output: Any = _NO_EVENT_OUTPUT, transport: dict | None = None) -> RunResult:
        run_id, logs = str(uuid.uuid4()), []
        started = datetime.now(timezone.utc)
        correlation_id = str(initial.get('correlationId') or initial.get('correlation_id') or run_id) if isinstance(initial, dict) else run_id
        execution_state = execution_state or {'activities': {}, 'tasks': {}}
        activity_outputs = execution_state.setdefault('activities', {})
        task_outputs = execution_state.setdefault('tasks', {})
        task_state = task_outputs.setdefault(process.id, {'name': process.name, 'activities': {}})
        context = {
            'input': initial, 'vars': {}, 'last': initial, 'resources': resources or {},
            'properties': properties or {}, 'project': project, 'runtime': self, 'logs': logs,
            'activities': activity_outputs, 'tasks': task_outputs,
            'context': {'taskId': process.id, 'activityId': '', 'environment': getattr(project, 'active_environment', '') if project else '', 'correlationId': correlation_id, 'runId': run_id},
            'transport': transport or {},
            '_process': process, 'groupStack': [], 'jdbcTransactions': {},
        }
        self.log(logs, 'INFO', f'Job started: {process.name}', kind='lifecycle', correlationId=correlation_id, runId=run_id, startedAt=log_timestamp(started))
        def finish(status: str, output: dict) -> RunResult:
            ended = datetime.now(timezone.utc); duration = round((ended - started).total_seconds() * 1000, 3)
            self.log(logs, 'INFO' if status == 'completed' else 'ERROR', f'Job {status}: {process.name} in {duration:.3f} ms', kind='lifecycle', correlationId=correlation_id, runId=run_id, endedAt=log_timestamp(ended), durationMs=duration)
            for entry in logs:
                entry.setdefault('correlationId', correlation_id); entry.setdefault('runId', run_id)
            return RunResult(run_id=run_id, correlation_id=correlation_id, started_at=log_timestamp(started), ended_at=log_timestamp(ended), duration_ms=duration, status=status, output=output, logs=logs, activity_outputs=activity_outputs, task_outputs=task_outputs)
        activity_by_id = {a.id: a for a in process.activities}
        try: group_plans = self.compile_groups(process)
        except Exception as exc:
            self.log(logs, 'ERROR', str(exc)); return finish('failed', {})
        incoming = {t.target for t in process.transitions}
        starts = [activity_by_id[entry_activity_id]] if entry_activity_id in activity_by_id else ([a for a in process.activities if a.type == 'start'] or [a for a in process.activities if a.id not in incoming and a.type != 'catch'])
        if len(starts) != 1:
            self.log(logs, 'ERROR', 'Process must have exactly one Start activity')
            return finish('failed', {})
        current = starts[0]
        triggered_event_pending = event_output is not _NO_EVENT_OUTPUT and current.id == entry_activity_id
        try:
            step_count = 0
            while True:
                step_count += 1
                if step_count > max(100000, len(process.activities) * 10000):
                    raise FabricFault('Execution step limit exceeded; check group loop conditions', fault_type='GROUP_ITERATION_LIMIT')
                entered = await self.enter_group_boundaries(current.id, context, group_plans)
                if entered is None: break
                if entered != current.id:
                    current = activity_by_id[entered]; continue
                activity_started = perf_counter()
                operation = str(current.config.get('operation') or current.type)
                self.log(logs, 'INFO', f'Activity started: {process.name} / {current.name}', kind='activity', taskId=process.id, runtimeActivityId=current.id, activityName=current.name, activityType=current.type, operation=operation)
                context['context']['activityId'] = current.id
                error = None
                try:
                    if triggered_event_pending:
                        context['last'] = event_output
                        triggered_event_pending = False
                        self.log(logs, 'INFO', f'Event delivered: {process.name} / {current.name}', kind='event', taskId=process.id, runtimeActivityId=current.id, activityName=current.name, activityType=current.type, operation=operation)
                    else:
                        context['last'] = await self.execute_with_policy(current, context)
                    self.record_activity_output(current, context['last'], context)
                except Exception as exc: error = exc
                activity_duration = round((perf_counter() - activity_started) * 1000, 3)
                if error:
                    self.log(logs, 'ERROR', f'Activity failed: {process.name} / {current.name} in {activity_duration:.3f} ms: {error}', kind='activity', taskId=process.id, runtimeActivityId=current.id, activityName=current.name, activityType=current.type, operation=operation, durationMs=activity_duration)
                else:
                    self.log(logs, 'INFO', f'Activity completed: {process.name} / {current.name} in {activity_duration:.3f} ms', kind='activity', taskId=process.id, runtimeActivityId=current.id, activityName=current.name, activityType=current.type, operation=operation, durationMs=activity_duration)
                outgoing = [t for t in process.transitions if t.source == current.id]
                if current.type in ('end', 'http_response'): break
                if error:
                    chosen = next((t for t in outgoing if t.type == 'error'), None)
                    fault = self.fault_payload(error, current.id)
                    context['last'] = fault; context['context']['error'] = fault; context['vars']['error'] = fault
                    if not chosen:
                        retry_target = await self.retry_failed_group(context, group_plans)
                        if retry_target:
                            current = activity_by_id[retry_target]
                            continue
                        while context.get('groupStack'):
                            failed_state = context['groupStack'].pop()
                            await self._finish_group(failed_state, group_plans[failed_state['id']], context, False)
                        used = set(context['context'].setdefault('handledCatchIds', []))
                        catches = [activity for activity in process.activities if activity.type == 'catch' and activity.id not in used]
                        caught = next((activity for activity in catches if self.as_bool(activity.config.get('catchAll', True)) or (activity.config.get('errorType') and activity.config.get('errorType') == fault['type']) or (activity.config.get('errorCode') and str(activity.config.get('errorCode')) == fault['code'])), None)
                        if not caught: raise error
                        context['context']['handledCatchIds'].append(caught.id)
                        self.log(logs, 'WARN', f'{caught.name} caught {fault["type"]}: {fault["message"]}', caught.id, kind='exception', fault=fault)
                        current = caught
                        continue
                else:
                    chosen_edges = self.select_group_transitions(self.eligible_success_transitions(outgoing, context), context, group_plans)
                if error:
                    chosen_edges = [chosen] if chosen else []
                if not chosen_edges:
                    target_id = await self.leave_group_boundaries(current.id, None, context, group_plans, success=not error)
                    if target_id:
                        current = activity_by_id[target_id]; continue
                    if context.get('groupStack'): raise RuntimeErrorWithLogs(f'{current.name} has no matching outgoing transition')
                    break
                for edge in chosen_edges:
                    self.log(logs, 'DEBUG', f'Transition selected: {current.name} -> {activity_by_id[edge.target].name} ({edge.type})', kind='trace', runtimeActivityId=current.id, transitionId=edge.id, parallel=len(chosen_edges) > 1)
                if len(chosen_edges) == 1:
                    target_id = await self.leave_group_boundaries(current.id, chosen_edges[0].target, context, group_plans, success=not error)
                    if target_id is None: break
                    current = activity_by_id[target_id]
                else:
                    # A task may have multiple eligible success edges.  They are
                    # a fan-out, not an ordered if/else chain. Run every branch
                    # and share execution state so all outputs remain mappable.
                    branch_initial = context['last']
                    if not isinstance(branch_initial, dict): branch_initial = {'payload': branch_initial}
                    branch_initial = {**branch_initial, 'correlationId': correlation_id}
                    results = await asyncio.gather(*(
                        self.run(process, branch_initial, resources=resources, properties=properties,
                                 entry_activity_id=edge.target, project=project,
                                 execution_state=execution_state, transport=transport)
                        for edge in chosen_edges
                    ))
                    for result in results:
                        logs.extend(result.logs)
                    failed = next((result for result in results if result.status != 'completed'), None)
                    if failed: raise RuntimeErrorWithLogs(f'Parallel branch failed from {current.name}')
                    final_output = results[-1].output if results else context['last']
                    task_state['output'] = final_output
                    return finish('completed', final_output)
            final_output = context['last'] if isinstance(context['last'], dict) else {'result': context['last']}
            task_state['output'] = final_output
            return finish('completed', final_output)
        except Exception as exc:
            while context.get('groupStack'):
                state = context['groupStack'].pop()
                try: await self._finish_group(state, group_plans[state['id']], context, False)
                except Exception: pass
            self.log(logs, 'ERROR', str(exc), current.id)
            task_state['error'] = {'message': str(exc), 'activityId': current.id}
            return finish('failed', {})

    def eligible_success_transitions(self, outgoing, context: dict):
        """Return every success edge eligible after an activity completes.

        The old implementation used ``next`` and therefore silently discarded
        every second success edge. Conditional edges retain their precedence;
        when one or more conditions match, all matching conditions are fanned
        out. Otherwise all normal success edges (or all success_no_match edges)
        are returned.
        """
        conditional = [edge for edge in outgoing if edge.type == 'success_condition' and self.condition(edge.condition, context)]
        if conditional: return conditional
        success = [edge for edge in outgoing if edge.type == 'success']
        return success or [edge for edge in outgoing if edge.type == 'success_no_match']

    @staticmethod
    def select_group_transitions(edges, context: dict, plans: dict[str, dict]):
        """Apply Pick First semantics before the generic parallel fan-out.

        Pick First is a first-match branch construct: once execution reaches a
        branch in the group, only that eligible branch continues. Ordinary
        success fan-out remains unchanged for every other group type.
        """
        if len(edges) <= 1:
            return edges
        active = next((state for state in reversed(context.get('groupStack', []))
                       if getattr(plans.get(state.get('id'), {}).get('group'), 'type', None) == 'pick_first'), None)
        if active:
            return edges[:1]
        process = context.get('_process')
        if not process:
            return edges
        pick_first_members = set().union(*(plan['members'] for plan in plans.values() if plan['group'].type == 'pick_first'))
        candidates = [edge for edge in edges if edge.target in pick_first_members]
        return edges[:1] if candidates else edges

    @staticmethod
    def record_activity_output(activity: Activity, result, ctx: dict):
        """Retain every executed activity result for downstream mappings and debugging."""
        record = {'activityId': activity.id, 'name': activity.name, 'type': activity.type, 'output': result}
        if activity.type in ('xml', 'flat') and activity.config.get('operation') == 'parse' and isinstance(result, dict) and result.get('xml'):
            record['displayOutput'] = result['xml']
        record.update(ctx.pop('_activityMetadata', {}))
        ctx.setdefault('activities', {})[activity.id] = record
        task_id = ctx.get('context', {}).get('taskId')
        if task_id:
            task = ctx.setdefault('tasks', {}).setdefault(task_id, {'activities': {}})
            task.setdefault('activities', {})[activity.id] = record

    @staticmethod
    def fault_payload(error: Exception, activity_id: str) -> dict:
        stack_trace = ''.join(traceback.format_exception(type(error), error, error.__traceback__))
        return {'type': getattr(error, 'fault_type', error.__class__.__name__), 'code': str(getattr(error, 'code', '') or ''), 'message': str(error), 'stackTrace': stack_trace, 'activityId': activity_id, 'details': getattr(error, 'details', {}) or {}, 'cause': getattr(error, 'cause', None)}

    def log(self, logs, level, message, activity_id=None, **details):
        logs.append({'time': log_timestamp(), 'level': level, 'message': message, 'activityId': activity_id, **details})

    @staticmethod
    def _cron_field_matches(expression: str, value: int, minimum: int, maximum: int) -> bool:
        def item_matches(item: str) -> bool:
            base, _, step_text = item.partition('/')
            step = int(step_text or 1)
            if step < 1: raise ValueError('Cron step must be greater than zero')
            if base == '*': start, end = minimum, maximum
            elif '-' in base: start, end = (int(part) for part in base.split('-', 1))
            else: start = end = int(base)
            if start < minimum or end > maximum or start > end: raise ValueError(f'Cron value {item!r} is outside {minimum}-{maximum}')
            return start <= value <= end and (value - start) % step == 0
        return any(item_matches(item.strip()) for item in expression.split(',') if item.strip())

    @classmethod
    def _next_cron_time(cls, expression: str, now: datetime) -> datetime:
        parts = expression.split()
        if len(parts) != 5: raise FabricFault('Cron expression requires five fields: minute hour day month weekday', fault_type='SCHEDULER')
        candidate = now.replace(second=0, microsecond=0) + timedelta(minutes=1)
        for _ in range(527040):
            weekday = (candidate.weekday() + 1) % 7
            try:
                if (cls._cron_field_matches(parts[0], candidate.minute, 0, 59) and cls._cron_field_matches(parts[1], candidate.hour, 0, 23) and cls._cron_field_matches(parts[2], candidate.day, 1, 31) and cls._cron_field_matches(parts[3], candidate.month, 1, 12) and cls._cron_field_matches(parts[4], weekday, 0, 6)):
                    return candidate
            except ValueError as error:
                raise FabricFault(str(error), fault_type='SCHEDULER') from error
            candidate += timedelta(minutes=1)
        raise FabricFault('Cron expression did not produce an execution time within one year', fault_type='SCHEDULER')

    async def execute_with_policy(self, activity: Activity, ctx: dict):
        policy = self.resolve(activity.config.get('errorPolicy', {}), ctx)
        advanced = self.resolve({
            'logPayload': '${properties.advanced.logPayload}',
            'retryEnabled': '${properties.advanced.retryEnabled}',
            'retryCount': '${properties.advanced.retryCount}',
            'retryIntervalSeconds': '${properties.advanced.retryIntervalSeconds}',
            **activity.config.get('advanced', {})
        }, ctx)
        log_payload = self.as_bool(advanced.get('logPayload', False))
        outbound_retry = self.is_outbound(activity) and self.as_bool(advanced.get('retryEnabled', False))
        legacy_retry = policy.get('action') == 'retry'
        advanced_count = advanced.get('retryCount')
        advanced_interval = advanced.get('retryIntervalSeconds')
        retry_count = int(3 if advanced_count in (None, '') else advanced_count) if outbound_retry else int(policy.get('retryCount', 0) or 0)
        attempts = retry_count + 1 if outbound_retry or legacy_retry else 1
        retry_delay = float(60 if advanced_interval in (None, '') else advanced_interval) if outbound_retry else float(policy.get('retryDelay', 0) or 0) / 1000
        if log_payload:
            self.log(ctx.setdefault('logs', []), 'INFO', f'{activity.name} input payload: {self.payload_text(ctx.get("last"))}', activity.id)
        for attempt in range(attempts):
            try:
                result = await self.execute(activity, ctx)
                if activity.config.get('outputName'): ctx['vars'][activity.config['outputName']] = result
                if log_payload:
                    self.log(ctx.setdefault('logs', []), 'INFO', f'{activity.name} output payload: {self.payload_text(result)}', activity.id)
                return result
            except Exception as exc:
                if attempt + 1 < attempts:
                    self.log(ctx.setdefault('logs', []), 'WARN', f'{activity.name} failed; retry {attempt + 1} of {retry_count} in {retry_delay:g} seconds: {exc}', activity.id)
                    if retry_delay: await asyncio.sleep(retry_delay)
                    continue
                fault = {'type': exc.__class__.__name__, 'description': str(exc), 'activityId': activity.id}
                if policy.get('includeInput', True): fault['input'] = ctx.get('last')
                if policy.get('outputVariable'): ctx['vars'][policy['outputVariable']] = fault
                if policy.get('action') == 'continue': return fault
                if policy.get('action') == 'ignore': return ctx.get('last')
                raise

    @staticmethod
    def as_bool(value):
        return value if isinstance(value, bool) else str(value).strip().lower() in ('true','1','yes','on')

    @staticmethod
    def payload_text(value):
        # IDoc listener/parser results carry the canonical wire representation
        # in IDocXML (or payload). Do not turn that XML into a JSON object just
        # because the surrounding application log is JSON-lines formatted.
        if isinstance(value, dict):
            for key in ('IDocXML', 'xmlString', 'xml', 'payload'):
                candidate = value.get(key)
                if isinstance(candidate, str) and candidate.lstrip().startswith('<'):
                    return candidate
        if isinstance(value, str) and value.lstrip().startswith('<'):
            return value
        try: return json.dumps(value, ensure_ascii=False, default=str)
        except (TypeError, ValueError): return str(value)

    @staticmethod
    def is_outbound(activity: Activity) -> bool:
        operation = activity.config.get('operation', '')
        if activity.type in ('http','jdbc','snowflake','amqp','ftp','sftp'): return True
        if activity.type == 'ems': return operation in ('send','publish','request_reply','reply')
        if activity.type == 'kafka': return operation in ('publish','get')
        if activity.type == 'pubsub': return operation == 'publish'
        if activity.type == 'rest': return operation == 'invoke'
        if activity.type == 'soap': return operation == 'request_reply'
        if activity.type == 'sap': return operation in ('idoc_acknowledgment','idoc_confirmation','post_idoc','invoke_rfc_bapi','reply_rfc_bapi','read_table')
        return False

    async def execute(self, activity: Activity, ctx: dict):
        # Resolve environment, input, variable, and previous-output expressions in every activity field.
        cfg = self.resolve(activity.config, ctx)
        for key, expression in activity.config.get('inputMappings', {}).items():
            include, value = self.evaluate_mapping(expression, ctx)
            if include: self.assign_path(cfg, key, value)
        if activity.type in ('ftp','sftp','http','http_listener','http_response','rest','soap','sap') and cfg.get('resourceId'):
            shared = ctx['resources'].get(cfg['resourceId'])
            if not shared: raise RuntimeError(f'{activity.name} requires a valid shared connection')
            cfg = {**self.resolve(shared.config, ctx), **cfg}
            if activity.type in ('http','rest','soap') and cfg.get('baseUrl') and cfg.get('url','').startswith('/'):
                cfg['url'] = cfg['baseUrl'].rstrip('/') + cfg['url']
        if activity.type == 'start':
            mapped = self.map_input_values(activity.config.get('inputMappings', {}), ctx)
            return self.unwrap_boundary(mapped, 'payload', ctx['input'])
        if activity.type == 'end':
            mapped = self.map_input_values(activity.config.get('inputMappings', {}), ctx)
            return self.unwrap_boundary(mapped, 'result', ctx['last'])
        if activity.type == 'timer':
            now = datetime.now(timezone.utc)
            environment = str(ctx.get('context', {}).get('environment') or '').lower()
            # The designer's Run once option is intended for local Run/Debug.
            # Debug sessions may execute against a named environment profile
            # (for example production), so identify Debug explicitly rather
            # than accidentally making it wait for the production schedule.
            is_debug = bool(ctx.get('context', {}).get('debugSessionId'))
            run_once = self.as_bool(cfg.get('runOnceOnLocalStart', True)) and (environment == 'local' or is_debug)
            mode = str(cfg.get('scheduleMode') or 'dateTime')
            if run_once:
                scheduled = now
                trigger_mode = 'local-run-once'
            elif mode == 'cron':
                scheduled = self._next_cron_time(str(cfg.get('cronExpression') or ''), now)
                trigger_mode = 'cron'
            else:
                raw = str(cfg.get('scheduledDateTime') or '').strip()
                if not raw: raise FabricFault('Scheduler requires a date/time or local Run once option', fault_type='SCHEDULER')
                try:
                    scheduled = datetime.fromisoformat(raw.replace('Z', '+00:00'))
                    if scheduled.tzinfo is None: scheduled = scheduled.astimezone()
                    scheduled = scheduled.astimezone(timezone.utc)
                except ValueError as error: raise FabricFault(f'Invalid scheduler date/time: {raw}', fault_type='SCHEDULER') from error
                trigger_mode = 'dateTime'
            delay = max(0.0, (scheduled - now).total_seconds())
            self.log(ctx['logs'], 'INFO', f'Scheduler armed for {scheduled.isoformat()}', activity.id, kind='scheduler', triggerMode=trigger_mode, waitSeconds=round(delay, 3))
            if delay: await asyncio.sleep(delay)
            fired = datetime.now(timezone.utc)
            return {'scheduledTime': scheduled.isoformat(), 'actualTime': fired.isoformat(), 'sequence': 1, 'triggerMode': trigger_mode, 'payload': ctx['last']}
        if activity.type == 'basic':
            operation = str(cfg.get('operation') or 'empty')
            if operation == 'empty': return ctx['last']
            if operation == 'assign':
                name, value = str(cfg.get('variable') or '').strip(), cfg.get('value', ctx['last'])
                if not name: raise RuntimeError('Assign Variable requires a process variable name')
                ctx['vars'][name] = value; return {'name': name, 'value': value}
            if operation == 'checkpoint':
                checkpoint_id, created = str(uuid.uuid4()), log_timestamp()
                checkpoint = {'checkpointId': checkpoint_id, 'name': str(cfg.get('checkpointName') or activity.name), 'timestamp': created, 'activityId': activity.id}
                if self.as_bool(cfg.get('includeProcessState', True)):
                    checkpoint['state'] = {'last': ctx.get('last'), 'vars': dict(ctx.get('vars', {})), 'context': dict(ctx.get('context', {}))}
                ctx.setdefault('checkpoints', []).append(checkpoint)
                self.log(ctx['logs'], 'INFO', f"Checkpoint created: {checkpoint['name']}", activity.id, kind='checkpoint', checkpointId=checkpoint_id)
                return {key: value for key, value in checkpoint.items() if key != 'state'}
            if operation == 'sleep':
                duration = float(cfg.get('duration') or 0); unit = str(cfg.get('unit') or 'milliseconds').lower()
                seconds = duration * (60 if unit == 'minutes' else 1 if unit == 'seconds' else .001)
                await asyncio.sleep(max(0, seconds)); return {'sleptMilliseconds': round(seconds * 1000), 'payload': ctx['last']}
            if operation == 'get_context': return dict(ctx.get('context', {}))
            if operation == 'set_context':
                values = cfg.get('values') or {}
                if not isinstance(values, dict): raise RuntimeError('Set Process Context requires an object')
                ctx['context'].update(values); return {'context': dict(ctx['context'])}
            if operation == 'get_shared_variable':
                name = str(cfg.get('name') or '').strip(); return {'name': name, 'value': self.shared_variables.get(name, cfg.get('default'))}
            if operation == 'set_shared_variable':
                name, value = str(cfg.get('name') or '').strip(), cfg.get('value', ctx['last'])
                if not name: raise RuntimeError('Set Shared Variable requires a name')
                self.shared_variables[name] = value; return {'name': name, 'value': value}
            if operation == 'external_command':
                command = str(cfg.get('command') or cfg.get('commandToExecute') or '').strip()
                if not command: raise FabricFault('External Command requires a command to execute', fault_type='InvalidInputException')
                arguments = shlex.split(command, posix=os.name != 'nt')
                if os.name == 'nt': arguments = [item[1:-1] if len(item) > 1 and item[0] == item[-1] and item[0] in ('"', "'") else item for item in arguments]
                if not arguments: raise FabricFault('External Command is empty', fault_type='InvalidInputException')
                environment = cfg.get('environment') or {}
                if isinstance(environment, str):
                    environment = dict(item.split('=', 1) for item in environment.split(',') if '=' in item)
                process_env = ({str(key): str(value) for key, value in environment.items()} if cfg.get('replaceEnvironment') else {**os.environ, **{str(key): str(value) for key, value in environment.items()}})
                try:
                    process = await asyncio.create_subprocess_exec(*arguments, cwd=cfg.get('workingDirectory') or None, env=process_env, stdin=asyncio.subprocess.PIPE, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE)
                    stdout, stderr = await asyncio.wait_for(process.communicate(str(cfg.get('input') or '').encode()), timeout=float(cfg.get('timeoutSeconds') or 300))
                except asyncio.TimeoutError as exc:
                    process.kill(); await process.wait(); raise FabricFault('External command exceeded its timeout', fault_type='CommandExecutionError') from exc
                except (OSError, ValueError) as exc: raise FabricFault(str(exc), fault_type='CommandExecutionError') from exc
                output, error = stdout.decode(cfg.get('encoding') or 'utf-8', errors='replace'), stderr.decode(cfg.get('encoding') or 'utf-8', errors='replace')
                output_file = str(cfg.get('outputFile') or cfg.get('outputFilename') or '').strip()
                if output_file:
                    try: Path(output_file).write_text(output + error, encoding=cfg.get('encoding') or 'utf-8')
                    except OSError as exc: raise FabricFault(str(exc), fault_type='FileIOError') from exc
                split = str(cfg.get('outputLineSplitting') or 'None')
                if split == 'AtOperatingSystemLineEnd': output, error = output.splitlines(), error.splitlines()
                elif split == 'AtSpecifiedToken':
                    token = str(cfg.get('splitToken') or '')
                    if not token: raise FabricFault('A split token is required', fault_type='InvalidInputException')
                    output, error = output.split(token), error.split(token)
                return {'returnCode': process.returncode, 'output': output if cfg.get('provideCommandOutput', True) else None, 'error': error if cfg.get('provideCommandOutput', True) else None, 'outputFile': output_file or None}
            raise RuntimeError(f'Unsupported Basic/General operation {operation}')
        if activity.type == 'catch': return ctx.get('context', {}).get('error') or ctx['last']
        if activity.type == 'throw':
            details = cfg.get('details', {})
            if isinstance(details, str) and details.strip():
                try: details = json.loads(details)
                except ValueError: details = {'text': details}
            if cfg.get('stackTrace'):
                details = {**(details if isinstance(details, dict) else {'value': details}), 'stackTrace': cfg.get('stackTrace')}
            raise FabricFault(str(cfg.get('message') or 'Business fault'), fault_type=str(cfg.get('errorType') or cfg.get('type') or 'UserDefinedException'), code=str(cfg.get('code') or ''), details=details)
        if activity.type == 'rethrow':
            fault = ctx.get('context', {}).get('error')
            if not fault: raise FabricFault('Rethrow requires an active caught exception', fault_type='RethrowException')
            raise FabricFault(fault.get('message', 'Rethrown exception'), fault_type=fault.get('type', 'RethrowException'), code=fault.get('code', ''), details=fault.get('details', {}), cause=fault)
        if activity.type in ('http_listener',) or (activity.type == 'rest' and cfg.get('operation') == 'receiver') or (activity.type == 'soap' and cfg.get('operation') == 'service'):
            return ctx['last']
        if activity.type == 'http_response':
            body = self.resolve(cfg.get('body', '${last}'), ctx)
            return {'__httpResponse': True, 'statusCode': int(cfg.get('statusCode', 200)), 'headers': cfg.get('headers', {}), 'body': body}
        if activity.type == 'log':
            level = str(cfg.get('level') or 'INFO').upper()
            if level not in ('DEBUG', 'INFO', 'WARN', 'ERROR'): level = 'INFO'
            message = cfg.get('message')
            if message in (None, ''): message = f'{activity.name} payload'
            if not isinstance(message, str): message = self.payload_text(message)
            payload = cfg.get('payload', ctx['last'])
            # Treat an untouched empty payload editor as "current activity
            # payload". This keeps a Log activity useful after an inbound
            # listener/parser without requiring a redundant mapping, while
            # still honoring an explicit payload mapping.
            if payload in (None, '') and 'payload' not in activity.config.get('inputMappings', {}):
                payload = ctx['last']
            # Existing projects commonly mapped a parser branch such as
            # `${activities.<id>.output.SAPIDoc.ARTMAS05}`. That branch is the
            # JSON view of the IDoc, while the parser also publishes IDocXML
            # as the canonical XML view. Preserve the old mapping and log the
            # requested IDoc in the configured output format.
            payload_mapping = activity.config.get('inputMappings', {}).get('payload')
            if isinstance(payload_mapping, str) and '.output.' in payload_mapping and isinstance(payload, dict):
                match = re.search(r'\$\{activities\.([^\.}]+)\.output(?:\.|\})', payload_mapping)
                source_record = ctx.get('activities', {}).get(match.group(1)) if match else None
                source_output = source_record.get('output') if isinstance(source_record, dict) else None
                if isinstance(source_output, dict) and str(source_output.get('format', '')).upper() == 'XML' and source_output.get('IDocXML'):
                    payload = source_output['IDocXML']
            include_payload = self.as_bool(cfg.get('includePayload', False)) or 'payload' in activity.config.get('inputMappings', {})
            event = {'level': level, 'message': message, 'activityId': activity.id, 'activityName': activity.name}
            if include_payload:
                event['payload'] = payload
                event['payloadFormat'] = 'XML' if isinstance(payload, str) and payload.lstrip().startswith('<') else 'JSON' if isinstance(payload, (dict, list)) else 'TEXT'
            self.log(ctx.setdefault('logs', []), level, message, activity.id, activityName=activity.name, **({'payload': payload, 'payloadFormat': event['payloadFormat']} if include_payload else {}))
            ctx['_activityMetadata'] = {'logEvent': event}
            return ctx['last']
        if activity.type == 'confirm':
            previous = ctx['last'] if isinstance(ctx['last'], dict) else {}
            handles = cfg.get('ackIds') or cfg.get('ackId') or previous.get('ackIds') or previous.get('ackId') or previous.get('AckID')
            if not handles:
                if self.as_bool(cfg.get('failIfMissing', True)):
                    raise RuntimeError('Confirm Message requires an acknowledgement handle from a client/manual receiver')
                return {'confirmed': False, 'count': 0, 'ackIds': [], 'technologies': []}
            return await self.confirm_messages(handles)
        if activity.type in ('mapper', 'transform', 'ai_transform'):
            source = {
                **(ctx['last'] if isinstance(ctx.get('last'), dict) else {}),
                'last': ctx.get('last'), 'input': ctx.get('input'),
                'properties': ctx.get('properties', {}), 'vars': ctx.get('vars', {}),
                'context': ctx.get('context', {}), 'activities': ctx.get('activities', {}),
            }
            rules = []
            for rule in activity.config.get('mappings', []) or []:
                normalized = {**rule}
                if 'constant' in rule: normalized['constant'] = self.resolve(rule['constant'], ctx)
                for key in ('source', 'select'):
                    value = normalized.get(key)
                    if isinstance(value, str) and value.startswith('${') and value.endswith('}'):
                        normalized[key] = value[2:-1]
                if isinstance(normalized.get('condition'), str):
                    normalized['condition'] = re.sub(r'\$\{([^}]+)\}', r'\1', normalized['condition'])
                if isinstance(normalized.get('whens'), list):
                    normalized['whens'] = [{
                        **branch,
                        'condition': re.sub(r'\$\{([^}]+)\}', r'\1', str(branch.get('condition') or '')),
                        'source': branch['source'][2:-1] if isinstance(branch.get('source'), str) and branch['source'].startswith('${') and branch['source'].endswith('}') else branch.get('source'),
                    } for branch in normalized['whens']]
                rules.append(normalized)
            return execute_mapping(source, rules, cfg)
        if activity.type == 'dataweave':
            try:
                transformed = execute_dataweave(
                    str(cfg.get('script') or '%dw 2.0\noutput application/json\n---\npayload'),
                    payload=cfg.get('payload', ctx.get('last')),
                    attributes=cfg.get('attributes', ctx.get('attributes', {})),
                    variables={**ctx.get('vars', {}), **(cfg.get('variables', {}) or {})},
                    input_mime_type=str(cfg.get('inputMimeType') or ''),
                )
                target = str(cfg.get('outputTarget') or 'payload').lower()
                if target == 'attributes': ctx['attributes'] = transformed
                elif target == 'variable':
                    name = str(cfg.get('outputVariable') or 'transformResult').strip()
                    if not name: raise DataWeaveError('A variable output target requires a variable name')
                    ctx.setdefault('vars', {})[name] = transformed
                return transformed
            except DataWeaveError as exc:
                raise FabricFault(str(exc), fault_type='DATAWEAVE_SYNTAX', cause=exc.__class__.__name__) from exc
        if activity.type == 'sap':
            resource = ctx['resources'].get(cfg.get('resourceId'))
            if not resource or resource.type != 'sap': raise RuntimeError('SAP activity requires an SAP ECC shared connection')
            sap_cfg = {**self.resolve(resource.config, ctx), **cfg}
            # SAP's TID manager is a separate shared resource in the reference
            # plug-in. Resolve it here so the listener uses the configured
            # durable store instead of silently falling back to a process-local
            # file. Keep the fallback for older projects that never selected a
            # TID resource.
            operation = cfg.get('operation', 'invoke_rfc_bapi')
            if operation == 'idoc_listener':
                tid_resource_id = cfg.get('tidManagerId')
                tid_resource = ctx['resources'].get(tid_resource_id) if tid_resource_id else None
                if tid_resource and tid_resource.type == 'sap_tid':
                    tid_cfg = self.resolve(tid_resource.config, ctx)
                    if str(tid_cfg.get('mode') or 'active').lower() not in ('none', 'disabled', 'off', 'false', '0'):
                        configured_store = tid_cfg.get('storageFile') or tid_cfg.get('tidStorePath') or tid_cfg.get('url')
                        if configured_store:
                            sap_cfg['tidStorePath'] = configured_store
                    else:
                        sap_cfg['tidManagement'] = 'disabled'
            selected_idoc = next((item for item in sap_cfg.get('idocCatalog', []) if item.get('idocType') == sap_cfg.get('idocType')), None) or sap_cfg.get('selectedIdoc')
            if selected_idoc:
                sap_cfg = {**sap_cfg, 'selectedIdoc': selected_idoc, 'idocType': sap_cfg.get('idocType') or selected_idoc.get('idocType'), 'extensionType': sap_cfg.get('extensionType') or selected_idoc.get('extensionType',''), 'release': sap_cfg.get('release') or selected_idoc.get('release',''), 'idocSchema': selected_idoc.get('schema')}
            payload = cfg.get('payload', ctx['last'])
            if operation in ('idoc_listener', 'rfc_bapi_listener'):
                source = str(sap_cfg.get('messagingSource') or 'NoMessaging').strip().lower().replace(' ', '')
                if operation == 'idoc_listener' and source not in ('', 'nomessaging', 'direct', 'sapjcorfc/idoc_inbound_asynchronous'):
                    technology = {'ems':'ems', 'jms':'jms', 'kafka':'kafka'}.get(source)
                    if not technology:
                        raise RuntimeError(f'Unsupported SAP IDoc messaging source {sap_cfg.get("messagingSource")!r}')
                    broker_cfg = {
                        **sap_cfg,
                        'resourceId': sap_cfg.get('messagingResourceId'),
                        'operation': 'queue_receiver' if technology == 'ems' else 'receive_message' if technology == 'jms' else 'receive',
                        'destination': sap_cfg.get('messagingDestination') or sap_cfg.get('destination') or sap_cfg.get('topic'),
                        'topic': sap_cfg.get('messagingDestination') or sap_cfg.get('topic'),
                        'maxMessages': 1,
                    }
                    received = await self.messaging(technology, broker_cfg, ctx)
                    messages = received.get('messages') if isinstance(received, dict) else []
                    first = messages[0] if isinstance(messages, list) and messages else {}
                    broker_payload = received.get('body') if technology in ('ems', 'jms') else first.get('data')
                    properties = received.get('properties', {}) if isinstance(received, dict) else {}
                    return {**received, 'payload': broker_payload, 'SAPIDoc': properties.get('SAPIDoc', {}) if isinstance(properties, dict) else {}, 'messagingSource': technology.upper()}
                return await sap_adapter.receive_idoc(sap_cfg)
            if operation == 'reply_rfc_bapi' and sap_adapter._mode(sap_cfg) != 'mock':
                delivery = ctx.get('transport') or {}
                delivery_id, listener_key = delivery.get('deliveryId'), delivery.get('listenerKey')
                if not delivery_id or not listener_key:
                    raise RuntimeError('Reply from RFC/BAPI must run in the flow started by an RFC/BAPI Listener request')
                response = payload if isinstance(payload, dict) else {'RESULT': payload}
                sap_adapter.acknowledge_idoc(listener_key, delivery_id, True, response)
                delivery['completed'] = True
                return {'replied': True, 'functionName': delivery.get('functionName'), 'response': response}
            return await asyncio.to_thread(sap_adapter.execute, operation, sap_cfg, payload)
        if activity.type == 'http':
            method, url = cfg.get('method','GET'), self.resolve(cfg.get('url',''), ctx)
            async with httpx.AsyncClient(timeout=float(cfg.get('timeout', 30))) as client:
                response = await client.request(method, url, headers=self.resolve(cfg.get('headers', {}), ctx), json=self.resolve(cfg.get('body'), ctx) or None)
                response.raise_for_status()
                try: return response.json()
                except ValueError: return {'statusCode': response.status_code, 'body': response.text}
        if activity.type == 'rest' and cfg.get('operation') == 'invoke':
            async with httpx.AsyncClient(timeout=float(cfg.get('timeout', 30))) as client:
                response = await client.request(cfg.get('method', 'GET'), cfg.get('url', ''), headers=cfg.get('headers', {}), params=cfg.get('query', {}), json=cfg.get('body') if cfg.get('bodyType', 'json') == 'json' else None, content=cfg.get('body') if cfg.get('bodyType') != 'json' else None)
                response.raise_for_status()
                try: payload = response.json()
                except ValueError: payload = response.text
                return {'statusCode': response.status_code, 'headers': dict(response.headers), 'body': payload}
        if activity.type == 'soap' and cfg.get('operation') == 'request_reply':
            envelope = cfg.get('envelope') or ctx['last']
            if not isinstance(envelope, (str, bytes)): envelope = self.render_xml(envelope, cfg.get('rootElement', 'Request'))
            headers = {'Content-Type': cfg.get('contentType', 'text/xml; charset=utf-8'), **cfg.get('headers', {})}
            if cfg.get('soapAction'): headers['SOAPAction'] = cfg['soapAction']
            async with httpx.AsyncClient(timeout=float(cfg.get('timeout', 30))) as client:
                response = await client.post(cfg.get('url', ''), content=envelope, headers=headers)
                response.raise_for_status(); return {'statusCode': response.status_code, 'headers': dict(response.headers), 'body': response.text}
        if activity.type == 'file':
            path = Path(self.resolve(cfg.get('path',''), ctx)).expanduser()
            operation = cfg.get('operation','read')
            def file_info(item: Path):
                stat = item.stat()
                return {'fullName':str(item.resolve()), 'fileName':item.name, 'location':str(item.parent.resolve()), 'configuredFileName':str(path), 'type':'directory' if item.is_dir() else 'file', 'readProtected':not os.access(item, os.R_OK), 'writeProtected':not os.access(item, os.W_OK), 'size':stat.st_size, 'lastModified':datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat()}
            if operation == 'write':
                if path.exists() and not self.as_bool(cfg.get('overwrite', False)) and not self.as_bool(cfg.get('append', False)): raise FileExistsError(f'{path} already exists')
                if self.as_bool(cfg.get('createDirectories', True)): path.parent.mkdir(parents=True, exist_ok=True)
                binary = str(cfg.get('writeAs', 'Text')).lower() == 'binary' or cfg.get('binaryContent') not in (None, '')
                content = cfg.get('binaryContent') if binary else cfg.get('textContent', cfg.get('content', ctx['last']))
                if binary:
                    raw = base64.b64decode(content) if isinstance(content, str) else bytes(content)
                    if str(cfg.get('compression', 'None')).lower() == 'gzip': raw = gzip.compress(raw)
                    with path.open('ab' if self.as_bool(cfg.get('append', False)) else 'wb') as stream: stream.write(raw)
                else:
                    text = content if isinstance(content, str) else json.dumps(content, indent=2)
                    if self.as_bool(cfg.get('addLineSeparator', False)): text += os.linesep
                    if str(cfg.get('compression', 'None')).lower() == 'gzip':
                        with gzip.open(path, 'at' if self.as_bool(cfg.get('append', False)) else 'wt', encoding=cfg.get('encoding') or 'utf-8') as stream: stream.write(text)
                    else:
                        with path.open('a' if self.as_bool(cfg.get('append', False)) else 'w', encoding=cfg.get('encoding') or 'utf-8') as stream: stream.write(text)
                return {'path': str(path), 'written': True, 'success': True, 'fileInfo': file_info(path)}
            if operation == 'list':
                pattern = cfg.get('pattern', '*'); matches = list(path.rglob(pattern) if self.as_bool(cfg.get('recursive', False)) else path.glob(pattern))
                list_type = str(cfg.get('listType', 'Files and Directories')).lower()
                if list_type == 'only files': matches = [item for item in matches if item.is_file()]
                if list_type == 'only directories': matches = [item for item in matches if item.is_dir()]
                sort_by = str(cfg.get('sortBy', 'Name')).lower(); key = (lambda item: item.stat().st_size) if sort_by == 'size' else ((lambda item: item.stat().st_mtime) if sort_by == 'last modified' else (lambda item: item.name.lower()))
                matches.sort(key=key, reverse=str(cfg.get('sortOrder', 'Ascending')).lower() == 'descending')
                return {'files': [file_info(item) for item in matches], 'count':len(matches)}
            if operation == 'delete':
                if not path.exists() and self.as_bool(cfg.get('ignoreMissing', False)): return {'path':str(path), 'deleted':False, 'success':True}
                if path.is_dir(): shutil.rmtree(path) if self.as_bool(cfg.get('recursive', False)) else path.rmdir()
                else: path.unlink()
                return {'path': str(path), 'deleted': True, 'success': True}
            if operation in ('rename', 'copy'):
                destination = Path(self.resolve(cfg.get('destination',''), ctx)).expanduser()
                if self.as_bool(cfg.get('createDirectories', False)): destination.parent.mkdir(parents=True, exist_ok=True)
                if destination.exists() and not self.as_bool(cfg.get('overwrite', False)): raise FileExistsError(f'{destination} already exists')
                if destination.exists() and self.as_bool(cfg.get('overwrite', False)):
                    shutil.rmtree(destination) if destination.is_dir() else destination.unlink()
                if operation == 'copy':
                    if path.is_dir(): shutil.copytree(path, destination, copy_function=shutil.copy2 if self.as_bool(cfg.get('preserveAttributes', True)) else shutil.copy)
                    else: (shutil.copy2 if self.as_bool(cfg.get('preserveAttributes', True)) else shutil.copy)(path, destination)
                else: path.rename(destination)
                return {'source': str(path), 'destination': str(destination), 'path':str(destination), 'operation': operation, 'success':True, 'fileInfo':file_info(destination)}
            if operation == 'poll':
                matches = list(path.rglob(cfg.get('pattern', '*')) if self.as_bool(cfg.get('recursive', False)) else path.glob(cfg.get('pattern', '*')))
                return {'files': [file_info(item) for item in matches], 'count': len(matches), 'eventType':cfg.get('eventType', 'Created')}
            info = file_info(path)
            if self.as_bool(cfg.get('excludeFileContent', False)): return {'path':str(path), 'fileInfo':info, **info}
            if str(cfg.get('readAs', 'Text')).lower() == 'binary': return {'path':str(path), 'binaryContent':base64.b64encode(path.read_bytes()).decode(), 'content':None, 'fileInfo':info, **info}
            content = path.read_text(encoding=cfg.get('encoding') or 'utf-8')
            return {'path': str(path), 'content': content, 'textContent':content, 'fileInfo':info, **info}
        if activity.type == 'jdbc':
            resource = ctx['resources'].get(cfg.get('resourceId'))
            if not resource or resource.type != 'jdbc': raise FabricFault('JDBC activity requires a valid shared JDBC connection', fault_type='JDBCConnectionNotFoundException')
            transaction = ctx.get('jdbcTransactions', {}).get(cfg.get('resourceId'))
            try: return await asyncio.to_thread(jdbc_adapter.execute, self.resolve(resource.config, ctx), cfg, transaction, transaction is None)
            except Exception as exc: raise FabricFault(str(exc), fault_type=getattr(exc, 'fault_type', 'JDBCSQLException')) from exc
        if activity.type == 'snowflake':
            resource = ctx['resources'].get(cfg.get('resourceId'))
            if not resource or resource.type != 'snowflake': raise FabricFault('Snowflake activity requires a valid Snowflake JDBC shared connection', fault_type='SNOWFLAKE_CONNECTION')
            try: return await asyncio.to_thread(snowflake_adapter.execute, self.resolve(resource.config, ctx), cfg, ctx.get('last'))
            except Exception as exc: raise FabricFault(str(exc), fault_type='SNOWFLAKE_DATABASE_JDBC', code=str(getattr(exc, 'code', '') or '500009')) from exc
        if activity.type == 'amqp':
            resource = ctx['resources'].get(cfg.get('resourceId'))
            if not resource or resource.type != 'amqp': raise FabricFault('AMQP activity requires a valid AMQP shared connection', fault_type='AMQPConnectionException')
            rcfg, operation = self.resolve(resource.config, ctx), str(cfg.get('operation') or 'get')
            destination = str(cfg.get('queueName') or cfg.get('topicName') or cfg.get('entityName') or rcfg.get('entityName') or 'default')
            if rcfg.get('mode') == 'memory':
                key = f'amqp:{destination}'
                if operation == 'send':
                    message_id = str(cfg.get('messageID') or uuid.uuid4()); body = cfg.get('body', cfg.get('message', ctx['last']))
                    self.messages.setdefault(key, []).append({'messageId': message_id, 'body': body, 'UserProperties': cfg.get('userProperties') or {}, 'MessageProperties': {'deliveryMode': cfg.get('deliveryMode', 'Persistent') == 'Persistent', 'messageID': message_id, 'expiration': cfg.get('expiration', 0), 'priority': cfg.get('priority', 4), 'type': cfg.get('type'), 'contentType': cfg.get('contentType'), 'correlationID': cfg.get('correlationID')}})
                    return {'sendResult': True, 'MessageId': message_id}
                if operation in ('get', 'receive'):
                    items = self.messages.setdefault(key, []); message = items.pop(0) if items else None
                    if not message: return {'received': False, 'body': None, 'UserProperties': {}, 'MessageProperties': {}}
                    if str(cfg.get('acknowledgeMode') or 'Auto') != 'Auto':
                        token = self.register_acknowledgement('amqp', message['messageId']); self.acknowledgements[token]['message'] = message; self.acknowledgements[token]['destination'] = key
                        message['settlementToken'] = token; message['ackId'] = token
                    return {'received': True, **message}
                if operation == 'dead_letter':
                    token = str(cfg.get('settlementToken') or ''); pending = self.acknowledgements.pop(token, None)
                    if not pending: raise FabricFault('Settlement token was not found or expired', fault_type='AMQPPluginException')
                    dead_key = f"{pending.get('destination', key)}:$deadletter"; self.messages.setdefault(dead_key, []).append({**pending.get('message', {}), 'deadLetterReason': cfg.get('deadLetterReason'), 'deadLetterErrorDescription': cfg.get('deadLetterErrorDescription')})
                    return {'status': 'Success', 'settlementToken': token, 'messageId': pending.get('messageId')}
            try:
                if operation == 'send': return await asyncio.to_thread(amqp_adapter.send, rcfg, cfg)
                if operation in ('get', 'receive'):
                    message, callback = await asyncio.to_thread(amqp_adapter.get, rcfg, cfg)
                    if not message: return {'received': False, 'body': None, 'UserProperties': {}, 'MessageProperties': {}}
                    if callback:
                        token = self.register_acknowledgement('amqp', message.get('messageId', ''), callback); message['settlementToken'] = token; message['ackId'] = token
                    return {'received': True, **message}
                if operation == 'dead_letter':
                    token = str(cfg.get('settlementToken') or ''); pending = self.acknowledgements.pop(token, None)
                    if not pending: raise FabricFault('Settlement token was not found or expired', fault_type='AMQPPluginException')
                    result = pending.get('callback')(True) if pending.get('callback') else None
                    if asyncio.iscoroutine(result): await result
                    return {'status': 'Success', 'settlementToken': token, 'messageId': pending.get('messageId')}
            except FabricFault: raise
            except Exception as exc: raise FabricFault(str(exc), fault_type=getattr(exc, 'fault_type', 'AMQPPluginException')) from exc
        if activity.type == 'excel':
            try: return await asyncio.to_thread(self.read_excel, cfg)
            except Exception as exc: raise FabricFault(str(exc), fault_type='EXCEL_READ') from exc
        if activity.type == 'ftp': return await asyncio.to_thread(self.ftp, cfg, ctx)
        if activity.type == 'sftp': return await asyncio.to_thread(self.sftp, cfg, ctx)
        if activity.type == 'xml':
            try:
                if cfg.get('operation') == 'parse': return self.parse_xml_activity(cfg, ctx)
                return self.render_xml_activity(cfg, ctx)
            except FabricFault: raise
            except (ValueError, TypeError, UnicodeError, LookupError) as exc: raise FabricFault(str(exc), fault_type='XMLParseException' if cfg.get('operation') == 'parse' else 'XMLRenderException', cause=exc.__class__.__name__) from exc
        if activity.type == 'json':
            try: return self.json_activity(cfg, ctx)
            except FabricFault: raise
            except (ValueError, TypeError, UnicodeError) as exc: raise FabricFault(str(exc), fault_type='JSONParserException' if cfg.get('operation') == 'parse' else 'JSONRenderException', cause=exc.__class__.__name__) from exc
        if activity.type == 'flat':
            try: return self.flat_data(cfg, ctx)
            except FabricFault: raise
            except (ValueError, TypeError, csv.Error) as exc: raise FabricFault(str(exc), fault_type='ParseDataException' if cfg.get('operation') == 'parse' else 'RenderDataException', cause=exc.__class__.__name__) from exc
        if activity.type == 'call_task':
            project = ctx.get('project')
            if not project: raise RuntimeError('Call Sub Task requires project execution context')
            dynamic_value = self.resolve(cfg.get('dynamicTaskId'), ctx) if cfg.get('dynamicTaskId') not in (None, '') else None
            task_id = str(dynamic_value).strip() if dynamic_value not in (None, '') else str(cfg.get('taskId') or '').strip()
            task = next((item for item in project.tasks if item.id == task_id and item.kind == 'subtask'), None)
            if not task:
                task = next((item for item in project.tasks if item.name.casefold() == task_id.casefold() and item.kind == 'subtask'), None)
            if not task or task.kind != 'subtask': raise RuntimeError(f'Sub Task {task_id!r} was not found')
            mapped_values = self.map_input_values(activity.config.get('inputMappings', {}), ctx)
            mapped = self.unwrap_boundary(mapped_values, 'payload', ctx['last'] if isinstance(ctx['last'], dict) else {'value':ctx['last']})
            execution_state = {'activities': ctx.setdefault('activities', {}), 'tasks': ctx.setdefault('tasks', {})}
            invocation = self.run(task, mapped, ctx['resources'], ctx['properties'], project=project, execution_state=execution_state)
            if cfg.get('spawn'):
                asyncio.create_task(invocation)
                return {'spawned': True, 'taskId': task.id}
            result = await invocation
            ctx.setdefault('logs', []).extend(result.logs)
            if result.status == 'failed': raise RuntimeError(result.logs[-1]['message'] if result.logs else f'Sub Task {task.name} failed')
            return result.output
        if activity.type in ('ems', 'jms', 'kafka', 'pubsub'):
            return await self.messaging(activity.type, cfg, ctx)
        if activity.type == 'java': return await self.java_worker(cfg, ctx['last'])
        if activity.type == 'python': return await self.python_worker(cfg, ctx['last'])
        raise RuntimeError(f'Unsupported activity type {activity.type}')

    async def messaging(self, technology: str, cfg: dict, ctx: dict):
        resource = ctx['resources'].get(cfg.get('resourceId'))
        if not resource or resource.type != technology:
            raise RuntimeError(f'{technology.upper()} activity requires a shared {technology.upper()} connection')
        rcfg = self.resolve(resource.config, ctx)
        operation = cfg.get('operation', 'publish')
        destination = cfg.get('destination') or cfg.get('queue') or cfg.get('topic') or cfg.get('subscription') or 'default'
        broker_key = f'{technology}:{resource.id}:{destination}'
        payload = self.resolve(cfg.get('data', cfg.get('message', '${last}')), ctx)
        def mapping(value):
            if isinstance(value, dict): return value
            if isinstance(value, str) and value.strip():
                try: return json.loads(value)
                except ValueError: return {}
            return {}
        def kafka_bytes(value, serializer: str):
            kind = str(serializer or 'String').lower()
            if value is None: return None
            if kind == 'json': return json.dumps(value, separators=(',', ':')).encode()
            if kind == 'byte array':
                if isinstance(value, bytes): return value
                try: return base64.b64decode(value, validate=True) if isinstance(value, str) else bytes(value)
                except (ValueError, TypeError): return str(value).encode()
            if kind == 'avro schema':
                raise FabricFault('Avro Schema serialization requires the configured Schema Registry runtime', fault_type='KafkaSchemaRegistryException')
            return value if isinstance(value, bytes) else str(value).encode()
        def kafka_value(raw, deserializer: str):
            if raw is None: return None
            kind = str(deserializer or 'String').lower()
            if kind == 'byte array': return base64.b64encode(raw).decode()
            if kind == 'json': return json.loads(raw.decode())
            if kind == 'avro schema': raise FabricFault('Avro Schema deserialization requires the configured Schema Registry runtime', fault_type='KafkaSchemaRegistryException')
            return raw.decode(errors='replace')
        attributes = {str(key): str(value) for key, value in {**mapping(cfg.get('dynamicProperties')), **mapping(cfg.get('attributes')), **mapping(cfg.get('headers'))}.items()}
        timestamp = log_timestamp()
        envelope = {'id': str(uuid.uuid4()), 'data': payload, 'attributes': attributes, 'destination': destination, 'technology': technology, 'timestamp': timestamp, 'key': cfg.get('key'), 'correlationId': cfg.get('correlationId'), 'replyTo': cfg.get('replyTo'), 'type': cfg.get('type'), 'priority': int(cfg.get('priority', 4) or 4), 'deliveryMode': cfg.get('deliveryMode', 'Persistent'), 'expiration': int(cfg.get('expiration', 0) or 0)}
        receive_ops = {'receive', 'subscribe', 'get', 'queue_receiver', 'topic_subscriber', 'get_queue_message', 'receive_message', 'wait_request'}
        client_ack = str(cfg.get('acknowledgeMode', '')).lower() in ('client', 'manual', 'explicit client', 'explicit client dups ok') or cfg.get('acknowledge') is False
        # Existing projects that explicitly selected the local broker retain
        # that behavior. New shared connections omit mode and use the real
        # configured provider.
        if rcfg.get('mode') == 'memory':
            queue = self.messages.setdefault(broker_key, [])
            if operation in receive_ops:
                maximum = int(cfg.get('maxMessages', 1))
                received, queue[:] = queue[:maximum], queue[maximum:]
                for item in received:
                    if client_ack: item['ackId'] = self.register_acknowledgement(technology, item['id'])
                if technology in ('ems', 'jms'):
                    first = received[0] if received else {}
                    return {'body': first.get('data'), 'headers': {'JMSMessageID': first.get('id'), 'JMSTimestamp': first.get('timestamp'), 'JMSCorrelationID': first.get('correlationId'), 'JMSReplyTo': first.get('replyTo'), 'JMSType': first.get('type'), 'JMSPriority': first.get('priority'), 'JMSDeliveryMode': first.get('deliveryMode'), 'JMSExpiration': first.get('expiration'), 'JMSRedelivered': False}, 'properties': first.get('attributes', {}), 'ackId': first.get('ackId'), 'messages': received, 'count': len(received)}
                if technology == 'pubsub':
                    first = received[0] if received else {}
                    return {'MessageID': first.get('id'), 'PublishTime': first.get('timestamp'), 'Data': first.get('data'), 'Attributes': first.get('attributes', {}), 'AckID': first.get('ackId'), 'ackId': first.get('ackId'), 'messages': received, 'count': len(received)}
                return {'messages': received, 'count': len(received), 'destination': destination, 'ackId': received[0].get('ackId') if received else None, 'ackIds': [item['ackId'] for item in received if item.get('ackId')]}
            queue.append(envelope)
            if technology == 'pubsub': return {'TopicName': destination, 'MessageID': envelope['id'], 'messageId': envelope['id'], 'published': True}
            if technology in ('ems', 'jms'): return {'messageId': envelope['id'], 'destination': destination, 'timestamp': envelope['timestamp'], 'published': True}
            envelope.update({'topic': destination, 'partition': int(cfg.get('partitionId', 0) or 0), 'offset': len(queue)-1})
            return {'messageId': envelope['id'], 'topic': destination, 'partition': envelope['partition'], 'offset': envelope['offset'], 'timestamp': envelope['timestamp'], 'published': True}
        if technology in ('ems', 'jms'):
            required = {'JMS connection URL':rcfg.get('serverUrl'), 'Username':rcfg.get('username'), 'Password':rcfg.get('password')}
            if str(rcfg.get('connectionFactoryType', 'Direct')).lower() == 'jndi':
                required.update({'JNDI context factory':rcfg.get('jndiContextFactory'), 'JNDI provider URL':rcfg.get('jndiProviderUrl'), 'JNDI username':rcfg.get('jndiUsername'), 'JNDI password':rcfg.get('jndiPassword'), 'JNDI connection factory':rcfg.get('connectionFactory')})
            missing = [name for name, value in required.items() if not str(value or '').strip()]
            if missing: raise FabricFault(f"Required connection values are missing: {', '.join(missing)}", fault_type='JMSConnectionException')
            options = {**cfg, 'topic': 'topic' in operation or operation in ('publish', 'topic_subscriber'), 'clientAcknowledge': client_ack, 'properties': attributes}
            try:
                if operation in receive_ops:
                    output = await asyncio.to_thread(execute_jms, rcfg, 'receive', str(destination), None, options)
                    if not output.get('received'):
                        return {'body': None, 'headers': {}, 'properties': {}, 'ackId': None, 'messages': [], 'count': 0}
                    body = output.get('body')
                    try: body = json.loads(body) if isinstance(body, str) else body
                    except ValueError: pass
                    message_id = output.get('headers', {}).get('JMSMessageID') or str(uuid.uuid4())
                    # The native bridge completes its provider receive before the
                    # short-lived Java process exits. Preserve the client-ack
                    # orchestration contract with a one-time confirmation token
                    # so Confirm Message can be mapped consistently for native
                    # EMS/JMS and the in-process providers.
                    ack_id = self.register_acknowledgement(technology, message_id) if client_ack else None
                    return {**output, 'body': body, 'ackId': ack_id, 'ackIds': [ack_id] if ack_id else [], 'messages': [{'id':message_id, 'data':body, 'attributes':output.get('properties', {}), 'ackId':ack_id}], 'count': 1}
                output = await asyncio.to_thread(execute_jms, rcfg, 'send', str(destination), payload, options)
                return {**output, 'destination':destination, 'timestamp':timestamp, 'published':True}
            except JavaBridgeError as exc:
                raise FabricFault(str(exc), fault_type='JMSConnectionException') from exc
        if technology == 'kafka':
            try: from confluent_kafka import Consumer, Producer, TopicPartition
            except ImportError: raise RuntimeError('External Kafka mode requires confluent-kafka')
            common = {'bootstrap.servers': rcfg['bootstrapServers'], 'client.id': rcfg.get('clientId') or f'integration-fabric-{uuid.uuid4()}', 'request.timeout.ms': int(rcfg.get('requestTimeoutMilliseconds', 30000) or 30000), 'reconnect.backoff.ms': int(rcfg.get('reconnectBackoffMilliseconds', 50) or 50), 'retry.backoff.ms': int(rcfg.get('retryBackoffMilliseconds', 100) or 100), **mapping(rcfg.get('clientProperties'))}
            if rcfg.get('securityProtocol'): common['security.protocol'] = rcfg['securityProtocol']
            sasl_mechanism = rcfg.get('saslMechanism') or ('PLAIN' if str(rcfg.get('authenticationType') or '').strip().lower() == 'api key / secret' else '')
            if sasl_mechanism: common['sasl.mechanism'] = sasl_mechanism
            if rcfg.get('username'): common['sasl.username'] = rcfg['username']
            if rcfg.get('password'): common['sasl.password'] = rcfg['password']
            if rcfg.get('sslCaLocation'): common['ssl.ca.location'] = rcfg['sslCaLocation']
            if rcfg.get('sslCertificateLocation'): common['ssl.certificate.location'] = rcfg['sslCertificateLocation']
            if rcfg.get('sslKeyLocation'): common['ssl.key.location'] = rcfg['sslKeyLocation']
            if rcfg.get('sslKeyPassword'): common['ssl.key.password'] = rcfg['sslKeyPassword']
            if rcfg.get('principalName'): common['sasl.kerberos.principal'] = rcfg['principalName']
            if operation in receive_ops:
                consumer_cfg = {**common, 'group.id': cfg.get('groupId') or rcfg.get('groupId', 'integration-fabric'), 'auto.offset.reset': cfg.get('autoOffsetReset', cfg.get('offsetReset', 'earliest')), 'enable.auto.commit': bool(cfg.get('enableAutoCommit', not client_ack)), 'fetch.min.bytes': int(cfg.get('fetchMinBytes', 1) or 1), 'max.poll.records': int(cfg.get('maxPollRecords', cfg.get('maxMessages', 1)) or 1), 'session.timeout.ms': int(cfg.get('sessionTimeoutMs', 45000) or 45000), 'heartbeat.interval.ms': int(cfg.get('heartbeatIntervalMs', 3000) or 3000), **mapping(cfg.get('additionalProperties'))}
                consumer = Consumer(consumer_cfg); topics = [item.strip() for item in str(destination).split(';') if item.strip()]
                if self.as_bool(cfg.get('assignCustomPartition', False)):
                    partitions = []
                    for token in re.split(r'[,;]', str(cfg.get('partitionIds') or cfg.get('partitionId') or '0')):
                        token = token.strip()
                        if not token: continue
                        if '-' in token:
                            start, end = token.split('-', 1); partitions.extend(range(int(start), int(end) + 1))
                        else: partitions.append(int(token))
                    offset = int(cfg.get('customOffset')) if str(cfg.get('seekPosition', '')).lower() == 'custom' and cfg.get('customOffset') not in (None, '') else -2 if str(cfg.get('seekPosition', '')).lower() == 'beginning' else -1
                    consumer.assign([TopicPartition(topic, partition, offset) for topic in topics for partition in partitions])
                else: consumer.subscribe(topics)
                messages, native_messages = [], []
                for _ in range(int(cfg.get('maxMessages', 1))):
                    message = consumer.poll(float(cfg.get('timeout', 1)))
                    if message and not message.error():
                        data = kafka_value(message.value(), cfg.get('valueDeserializer', 'String'))
                        key = kafka_value(message.key(), cfg.get('keyDeserializer', 'String')) if message.key() else None
                        item = {'id':f'{message.topic()}:{message.partition()}:{message.offset()}','data':data,'key':key,'topic':message.topic(),'partition':message.partition(),'offset':message.offset(),'timestamp':message.timestamp()[1],'headers':dict(message.headers() or [])}
                        messages.append(item); native_messages.append(message)
                if client_ack:
                    pending = {'count': len(native_messages)}
                    def kafka_confirm(native_message):
                        consumer.commit(message=native_message, asynchronous=False)
                        pending['count'] -= 1
                        if pending['count'] <= 0: consumer.close()
                    for item, native_message in zip(messages, native_messages):
                        item['ackId'] = self.register_acknowledgement('kafka', item['id'], lambda m=native_message: kafka_confirm(m))
                    if not native_messages: consumer.close()
                else: consumer.close()
                return {'messages': messages, 'count': len(messages), 'ackId': messages[0].get('ackId') if messages else None, 'ackIds':[item['ackId'] for item in messages if item.get('ackId')]}
            producer_cfg = {**common, 'acks': str(cfg.get('acks','all')), 'compression.type': cfg.get('compressionType','none'), 'retries': int(cfg.get('retries',3) or 0), 'batch.size': int(cfg.get('batchSize',16384) or 16384), 'linger.ms': int(cfg.get('lingerMs',0) or 0), 'queue.buffering.max.kbytes': max(1, int(cfg.get('bufferMemory',33554432) or 33554432) // 1024), 'message.max.bytes': int(cfg.get('maxRequestSize',1048576) or 1048576), 'enable.idempotence': bool(cfg.get('enableIdempotence',False)), **mapping(cfg.get('additionalProperties'))}
            if cfg.get('transactionalId'): producer_cfg['transactional.id'] = cfg['transactionalId']
            producer = Producer(producer_cfg); delivered = {}; transactional = bool(cfg.get('transactionalId'))
            if transactional: producer.init_transactions(); producer.begin_transaction()
            def delivery(error, message):
                if error: delivered['error'] = str(error)
                else: delivered.update({'partition':message.partition(),'offset':message.offset(),'timestamp':message.timestamp()[1]})
            raw = kafka_bytes(payload, cfg.get('valueSerializer', 'String'))
            key = kafka_bytes(cfg.get('key'), cfg.get('keySerializer', 'String'))
            try:
                producer.produce(destination, raw, key=key, partition=int(cfg['partitionId']) if cfg.get('assignCustomPartition') else -1, headers=list(attributes.items()), callback=delivery); producer.flush()
                if transactional: producer.commit_transaction()
            except Exception:
                if transactional: producer.abort_transaction()
                raise
            if delivered.get('error'): raise RuntimeError(delivered['error'])
            return {**envelope, **delivered, 'messageId':envelope['id'], 'topic':destination, 'published':True}
        if technology == 'pubsub':
            try: from google.cloud import pubsub_v1
            except ImportError: raise RuntimeError('External Google Pub/Sub mode requires google-cloud-pubsub')
            client_kwargs, project_id = pubsub_client_configuration({**rcfg, 'projectId': cfg.get('projectId') or rcfg.get('projectId')})
            if operation in receive_ops:
                subscriber = create_pubsub_client(pubsub_v1.SubscriberClient, {**rcfg, 'projectId': cfg.get('projectId') or rcfg.get('projectId')}); path = subscriber.subscription_path(project_id, destination); response = subscriber.pull(request={'subscription': path, 'max_messages': int(cfg.get('maxMessages', 1))}, timeout=float(cfg.get('receiveTimeout', cfg.get('timeout', 10))))
                messages, native_ack_ids = [], []
                for item in response.received_messages:
                    record = {'id':item.message.message_id,'messageId':item.message.message_id,'data':item.message.data.decode(errors='replace'),'attributes':dict(item.message.attributes),'publishTime':item.message.publish_time.isoformat() if item.message.publish_time else None}
                    messages.append(record); native_ack_ids.append(item.ack_id)
                if client_ack:
                    pending = {'count': len(native_ack_ids)}
                    def pubsub_confirm(native_ack_id):
                        subscriber.acknowledge(request={'subscription':path,'ack_ids':[native_ack_id]})
                        pending['count'] -= 1
                        if pending['count'] <= 0: subscriber.close()
                    for record, native_ack_id in zip(messages, native_ack_ids):
                        record['ackId'] = self.register_acknowledgement('pubsub', record['id'], lambda ack=native_ack_id: pubsub_confirm(ack))
                    if not native_ack_ids: subscriber.close()
                else:
                    if response.received_messages: subscriber.acknowledge(request={'subscription':path,'ack_ids':[item.ack_id for item in response.received_messages]})
                    subscriber.close()
                first = messages[0] if messages else {}
                return {'MessageID':first.get('messageId'),'PublishTime':first.get('publishTime'),'Data':first.get('data'),'Attributes':first.get('attributes',{}),'AckID':first.get('ackId'),'ackId':first.get('ackId'),'messages':messages,'count':len(messages)}
            publisher = create_pubsub_client(pubsub_v1.PublisherClient, {**rcfg, 'projectId': cfg.get('projectId') or rcfg.get('projectId')})
            publish_timeout = max(1.0, float(cfg.get('publishTimeout', 60) or 60))
            try:
                path = publisher.topic_path(project_id, destination)
                raw = payload if isinstance(payload, bytes) else (payload.encode() if isinstance(payload, str) else json.dumps(payload).encode())
                try:
                    from google.api_core.retry import Retry
                    publish_future = publisher.publish(
                        path, raw, ordering_key=str(cfg.get('orderingKey', '')),
                        retry=Retry(deadline=publish_timeout),
                        timeout=publish_timeout, **attributes,
                    )
                    message_id = publish_future.result(timeout=publish_timeout + 2)
                except TimeoutError as exc:
                    raise RuntimeError(f'Google Pub/Sub publish timed out after {publish_timeout:g} seconds. Verify the topic, IAM permission, endpoint, proxy, and firewall settings.') from exc
                return {**envelope, 'TopicName': path, 'MessageID': message_id, 'messageId': message_id, 'published': True}
            finally:
                # Stop the batching/sequencer threads before closing gRPC.
                # Closing the channel first causes the Google batch thread to
                # raise "Cannot invoke RPC on closed channel" during retries.
                try: publisher.stop()
                except Exception: pass
                try: publisher.transport.close()
                except Exception: pass
        raise RuntimeError(f'Unsupported messaging technology {technology}')

    @staticmethod
    def assign_path(target: dict, path: str, value):
        parts = [part for part in str(path).split('.') if part]
        if not parts: return
        current = target
        for part in parts[:-1]:
            child = current.get(part)
            if not isinstance(child, dict): child = {}; current[part] = child
            current = child
        current[parts[-1]] = value

    def map_input_values(self, mappings: dict, ctx: dict) -> dict:
        result = {}
        for path, expression in (mappings or {}).items():
            include, value = self.evaluate_mapping(expression, ctx)
            if include: self.assign_path(result, path, value)
        return result

    def evaluate_mapping(self, expression, ctx) -> tuple[bool, Any]:
        """Evaluate the structured mapping statements used by every activity Input tab."""
        if not isinstance(expression, dict) or '$rule' not in expression:
            return True, self.resolve(expression, ctx)
        rule = str(expression.get('$rule', '')).lower()
        source_expression = expression.get('select') or expression.get('source')
        source = self.resolve(source_expression, ctx)
        if rule == 'if':
            return self.condition(expression.get('condition', ''), ctx), source
        if rule == 'when-otherwise':
            if self.condition(expression.get('condition', ''), ctx): return True, source
            return True, self.resolve(expression.get('otherwise'), ctx)
        if rule == 'choose':
            branches = expression.get('whens') or []
            if not branches and isinstance(expression.get('source'), str):
                try: branches = json.loads(expression['source'])
                except (TypeError, ValueError, json.JSONDecodeError): branches = []
            for branch in branches:
                if self.condition(branch.get('condition', ''), ctx): return True, self.resolve(branch.get('source'), ctx)
            return True, self.resolve(expression.get('otherwise'), ctx)
        if rule == 'for-each':
            values = source if isinstance(source, list) else ([] if source in (None, '') else [source])
            return True, values
        if rule == 'for-each-group':
            values = source if isinstance(source, list) else ([] if source in (None, '') else [source])
            group_path = str(expression.get('groupBy') or '').strip('.')
            groups: dict[str, list[Any]] = {}
            for item in values:
                current = item
                for part in group_path.split('.') if group_path else []:
                    current = current.get(part) if isinstance(current, dict) else None
                groups.setdefault(str(current), []).append(item)
            return True, [{'key': key, 'items': items} for key, items in groups.items()]
        return True, source

    @staticmethod
    def unwrap_boundary(mapped: dict, wrapper: str, fallback):
        if not mapped: return fallback
        return mapped[wrapper] if set(mapped) == {wrapper} else mapped

    @staticmethod
    def split_function_args(raw: str) -> list[str]:
        args, start, depth, quote = [], 0, 0, None
        for index, char in enumerate(raw):
            if quote:
                if char == quote and (index == 0 or raw[index - 1] != '\\'): quote = None
            elif char in ('"', "'"): quote = char
            elif char == '(': depth += 1
            elif char == ')': depth = max(0, depth - 1)
            elif char == ',' and depth == 0: args.append(raw[start:index].strip()); start = index + 1
        if raw.strip(): args.append(raw[start:].strip())
        return args

    def evaluate_function_expression(self, expression: str, ctx: dict, variables: dict | None = None):
        variables = variables or {}
        text = expression.strip()
        if text in variables: return variables[text]
        if len(text) >= 2 and text[0] == text[-1] and text[0] in ('"', "'"): return text[1:-1]
        if re.fullmatch(r'-?\d+(\.\d+)?', text): return float(text) if '.' in text else int(text)
        if text.lower() in ('true()', 'true'): return True
        if text.lower() in ('false()', 'false'): return False
        call = re.fullmatch(r'([\w:-]+)\((.*)\)', text, re.S)
        if not call: return self.resolve(text, ctx)
        name = call.group(1); args = [self.evaluate_function_expression(item, ctx, variables) for item in self.split_function_args(call.group(2))]
        if name.startswith('custom:'):
            function_name = name.split(':', 1)[1]
            project = ctx.get('project')
            definition = next((item for item in (getattr(project, 'custom_functions', []) if project else []) if item.name == function_name), None)
            if not definition: raise RuntimeError(f'Custom function {function_name!r} was not found in this project')
            bindings = {f'${parameter}': args[index] if index < len(args) else None for index, parameter in enumerate(definition.parameters)}
            return self.evaluate_function_expression(definition.expression, ctx, bindings)
        try: return apply_function(name, args[0] if args else None, args[1:])
        except (TypeError, ValueError, IndexError) as exc: raise RuntimeError(str(exc)) from exc

    def resolve(self, value, ctx):
        if isinstance(value, dict): return {key: self.resolve(item, ctx) for key, item in value.items()}
        if isinstance(value, list): return [self.resolve(item, ctx) for item in value]
        if isinstance(value, str) and value.startswith('__fabric_constant__:'):
            try: return json.loads(value.split(':', 1)[1])
            except (TypeError, ValueError, json.JSONDecodeError): return value
        if value == '${last}': return ctx['last']
        if value == '${input}': return ctx['input']
        if value == '${properties}': return ctx['properties']
        if value == '${vars}': return ctx['vars']
        if isinstance(value, str) and value.startswith('${last.') and value.endswith('}'):
            current = ctx['last']
            for part in value[7:-1].split('.'): current = current.get(part, '') if isinstance(current, dict) else ''
            return current
        if isinstance(value, str) and value.startswith('${input.'):
            current = ctx['input']
            for part in value[8:-1].split('.'): current = current.get(part, '') if isinstance(current, dict) else ''
            return current
        if isinstance(value, str) and value.startswith('${vars.'):
            current = ctx['vars']
            for part in value[7:-1].split('.'): current = current.get(part, '') if isinstance(current, dict) else ''
            return current
        if isinstance(value, str) and value.startswith('${properties.') and value.endswith('}'):
            return ctx['properties'].get(value[13:-1], '')
        if isinstance(value, str) and value.startswith('${') and value.endswith('}'):
            path = value[2:-1].split('.')
            if path[0] in ('activities', 'tasks', 'context'):
                current = ctx.get(path[0], {})
                for part in path[1:]:
                    if isinstance(current, dict): current = current.get(part, '')
                    elif isinstance(current, list) and part.isdigit() and int(part) < len(current): current = current[int(part)]
                    else: return ''
                return current
        if isinstance(value, str) and '${properties.' in value:
            return re.sub(r'\$\{properties\.([^}]+)\}', lambda match: str(ctx['properties'].get(match.group(1), '')), value)
        if isinstance(value, str) and re.fullmatch(r'[\w:-]+\(.*\)', value.strip(), re.S):
            return self.evaluate_function_expression(value, ctx)
        if isinstance(value, str) and len(value) >= 2 and value[0] == value[-1] and value[0] in ('"', "'"):
            try: return ast.literal_eval(value)
            except (SyntaxError, ValueError): return value[1:-1]
        return value

    def condition(self, expression, ctx):
        expression = str(expression or '').strip()
        def split_logical(value, keyword):
            depth = 0; quote = None; start = 0; parts = []; token = f' {keyword} '; index = 0
            while index < len(value):
                char = value[index]
                if quote:
                    if char == quote and (index == 0 or value[index-1] != '\\'): quote = None
                elif char in ('"', "'"): quote = char
                elif char == '(': depth += 1
                elif char == ')': depth = max(0, depth - 1)
                elif depth == 0 and value[index:index+len(token)].lower() == token:
                    parts.append(value[start:index].strip()); start = index + len(token); index = start - 1
                index += 1
            if parts: parts.append(value[start:].strip())
            return parts
        for keyword, evaluator in (('or', any), ('and', all)):
            parts = split_logical(expression, keyword)
            if parts: return evaluator(self.condition(part, ctx) for part in parts)
        if expression.lower().startswith('not(') and expression.endswith(')'): return not self.condition(expression[4:-1], ctx)
        function = re.fullmatch(r'(exists|empty|contains|startsWith|endsWith|matches)\((.*)\)', expression, re.I)
        if function:
            raw = function.group(2); args = []; depth = 0; quote = None; start = 0
            for index, char in enumerate(raw):
                if quote:
                    if char == quote and (index == 0 or raw[index-1] != '\\'): quote = None
                elif char in ('"', "'"): quote = char
                elif char == '(': depth += 1
                elif char == ')': depth = max(0, depth - 1)
                elif char == ',' and depth == 0: args.append(raw[start:index].strip()); start = index + 1
            args.append(raw[start:].strip())
            def value(arg):
                if len(arg) >= 2 and arg[0] == arg[-1] and arg[0] in ('"', "'"): return arg[1:-1]
                if re.fullmatch(r'-?\d+(\.\d+)?', arg): return float(arg) if '.' in arg else int(arg)
                return self.resolve(arg, ctx)
            values = [value(arg) for arg in args]
            name = function.group(1).lower()
            if name == 'exists': return bool(values and values[0] is not None and values[0] != '')
            if name == 'empty': return not values or values[0] is None or values[0] == '' or values[0] == [] or values[0] == {}
            if len(values) < 2: return False
            left, right = str(values[0]), str(values[1])
            if name == 'contains': return right in left
            if name == 'startswith': return left.startswith(right)
            if name == 'endswith': return left.endswith(right)
            if name == 'matches': return re.search(right, left) is not None
        for operator in ('==', '!=', '>=', '<=', '>', '<', '='):
            if operator in expression:
                left, right = (part.strip() for part in expression.split(operator, 1))
                def comparison_value(raw):
                    if len(raw) >= 2 and raw[0] == raw[-1] and raw[0] in ('"', "'"):
                        return raw[1:-1]
                    if re.fullmatch(r'-?\d+(\.\d+)?', raw):
                        return float(raw) if '.' in raw else int(raw)
                    return self.resolve(raw, ctx)
                left, right = comparison_value(left), comparison_value(right)
                try:
                    if operator in ('=', '=='): return left == right or str(left) == str(right)
                    if operator == '!=': return left != right and str(left) != str(right)
                    if operator == '>=': return left >= right
                    if operator == '<=': return left <= right
                    if operator == '>': return left > right
                    if operator == '<': return left < right
                except TypeError: return False
        if expression.lower() in ('true', 'true()', 'false', 'false()'): return expression.lower() in ('true', 'true()')
        return bool(self.resolve(expression, ctx))

    def jdbc(self, cfg, ctx):
        resource = ctx['resources'].get(cfg.get('resourceId'))
        if not resource: raise RuntimeError('JDBC activity requires a valid shared JDBC resource')
        rcfg, operation = self.resolve(resource.config, ctx), cfg.get('operation', 'query')
        if rcfg.get('driver', 'sqlite') != 'sqlite':
            raise RuntimeError(f"Driver {rcfg.get('driver')} requires its optional Python database adapter")
        conn = sqlite3.connect(self.resolve(rcfg.get('url', 'integration.db'), ctx)); conn.row_factory = sqlite3.Row
        try:
            parameters = {name: self.resolve(expression, ctx) for name, expression in cfg.get('parameters', {}).items()}
            sql = self.resolve(cfg.get('sql', ''), ctx)
            if operation == 'truncate':
                if sql.strip().lower().startswith('truncate '): sql = 'DELETE FROM ' + sql.strip().split()[-1]
            if operation == 'call': raise RuntimeError('Stored procedures are not supported by SQLite; select a server JDBC driver')
            cursor = conn.execute(sql, parameters)
            if operation in ('query', 'dynamic') and cursor.description:
                return {'rows': [dict(row) for row in cursor.fetchall()], 'rowCount': cursor.rowcount}
            conn.commit(); return {'rowCount': cursor.rowcount, 'lastInsertId': cursor.lastrowid}
        finally: conn.close()

    @staticmethod
    def read_excel(cfg: dict) -> dict:
        file_path = Path(str(cfg.get('filePath') or cfg.get('path') or ''))
        if not file_path.is_file(): raise FileNotFoundError(f'Excel workbook not found: {file_path}')
        def assign(target: dict, path: str, value):
            parts = [part.strip() for part in str(path).split('.') if part.strip()]
            current = target
            for part in parts[:-1]: current = current.setdefault(part, {})
            if parts: current[parts[-1]] = value
        if file_path.suffix.lower() == '.xls':
            try: import xlrd
            except ImportError as exc: raise RuntimeError('Legacy .xls workbooks require xlrd') from exc
            workbook = xlrd.open_workbook(file_path)
            requested = str(cfg.get('sheetName') or '').strip(); names = [requested] if requested else workbook.sheet_names()
            header_row = max(1, int(cfg.get('headerRow') or 1)) - 1; start_row = max(header_row + 1, int(cfg.get('startRow') or header_row + 2) - 1)
            maximum = max(0, int(cfg.get('maximumRows') or 0)); nested = bool(cfg.get('nestedHeaders', True)); sheets = []
            for name in names:
                sheet = workbook.sheet_by_name(name); headers = [str(sheet.cell_value(header_row, col)).strip() or f'column{col + 1}' for col in range(sheet.ncols)]; rows = []
                for row_index in range(start_row, sheet.nrows):
                    if maximum and len(rows) >= maximum: break
                    values = [sheet.cell_value(row_index, col) for col in range(sheet.ncols)]
                    if cfg.get('skipBlankRows', True) and all(value in (None, '') for value in values): continue
                    record = {}
                    for key, value in zip(headers, values): assign(record, key, value) if nested and '.' in key else record.__setitem__(key, value)
                    rows.append(record)
                sheets.append({'name': name, 'index': workbook.sheet_names().index(name), 'rowCount': len(rows), 'columnCount': len(headers), 'headers': headers, 'rows': rows})
            return {'workbook': {'fileName': file_path.name, 'sheetCount': len(sheets), 'sheets': sheets}, 'sheets': sheets}
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError('Excel activities require openpyxl') from exc
        workbook = load_workbook(file_path, read_only=True, data_only=bool(cfg.get('dataOnly', True)))
        try:
            requested = str(cfg.get('sheetName') or '').strip()
            names = [requested] if requested else list(workbook.sheetnames)
            header_row = max(1, int(cfg.get('headerRow') or 1)); start_row = max(header_row + 1, int(cfg.get('startRow') or header_row + 1))
            maximum = max(0, int(cfg.get('maximumRows') or 0)); nested = bool(cfg.get('nestedHeaders', True)); sheets = []
            for index, name in enumerate(names):
                if name not in workbook.sheetnames: raise ValueError(f'Worksheet {name!r} does not exist')
                sheet = workbook[name]
                headers = [str(cell.value).strip() if cell.value not in (None, '') else f'column{cell.column}' for cell in sheet[header_row]]
                rows = []
                for values in sheet.iter_rows(min_row=start_row, values_only=True):
                    if maximum and len(rows) >= maximum: break
                    if cfg.get('skipBlankRows', True) and all(value is None for value in values): continue
                    record = {}
                    for key, value in zip(headers, values):
                        if nested and '.' in key: assign(record, key, value)
                        else: record[key] = value
                    rows.append(record)
                sheets.append({'name': name, 'index': workbook.sheetnames.index(name), 'rowCount': len(rows), 'columnCount': len(headers), 'headers': headers, 'rows': rows})
            return {'workbook': {'fileName': file_path.name, 'sheetCount': len(sheets), 'sheets': sheets}, 'sheets': sheets}
        finally: workbook.close()

    def ftp(self, cfg, ctx):
        client_class = ftplib.FTP_TLS if cfg.get('tls') else ftplib.FTP
        client = client_class(); client.connect(cfg['host'], int(cfg.get('port', 21)), timeout=float(cfg.get('timeout', 30))); client.login(cfg.get('username', 'anonymous'), cfg.get('password', ''))
        if cfg.get('tls'): client.prot_p()
        try:
            if cfg.get('workingDirectory'): client.cwd(cfg['workingDirectory'])
            operation, remote = cfg.get('operation'), cfg.get('remotePath', '')
            if operation == 'get':
                buffer = io.BytesIO(); client.retrbinary(f'RETR {remote}', buffer.write); return {'remotePath': remote, 'contentBase64': __import__('base64').b64encode(buffer.getvalue()).decode(), 'size': buffer.tell()}
            if operation == 'put':
                content = cfg.get('content', ctx['last']); raw = content if isinstance(content, bytes) else (content.encode(cfg.get('encoding', 'utf-8')) if isinstance(content, str) else json.dumps(content).encode()); client.storbinary(f'STOR {remote}', io.BytesIO(raw)); return {'remotePath': remote, 'written': True, 'size': len(raw)}
            if operation == 'delete': client.delete(remote); return {'remotePath': remote, 'deleted': True}
            if operation == 'dir': return {'directory': client.pwd(), 'entries': client.nlst(remote) if remote else client.nlst()}
            if operation == 'change_dir': client.cwd(remote); return {'directory': client.pwd()}
            raise RuntimeError(f'Unsupported FTP operation {operation}')
        finally:
            try: client.quit()
            except Exception: client.close()

    def sftp(self, cfg, ctx):
        try: import paramiko
        except ImportError: raise RuntimeError('SFTP activities require the optional paramiko package')
        ssh = paramiko.SSHClient(); ssh.load_system_host_keys()
        if cfg.get('knownHostsFile'): ssh.load_host_keys(cfg['knownHostsFile'])
        ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy() if cfg.get('allowUnknownHostKey') else paramiko.RejectPolicy())
        ssh.connect(cfg['host'], port=int(cfg.get('port', 22)), username=cfg.get('username'), password=cfg.get('password') or None, key_filename=cfg.get('privateKeyFile') or None, timeout=float(cfg.get('timeout', 30)), allow_agent=bool(cfg.get('useSshAgent', False)), look_for_keys=bool(cfg.get('useSshAgent', False)))
        client = ssh.open_sftp()
        try:
            if cfg.get('workingDirectory'): client.chdir(cfg['workingDirectory'])
            operation, remote = cfg.get('operation'), cfg.get('remotePath', '')
            if operation == 'get':
                buffer = io.BytesIO(); client.getfo(remote, buffer); return {'remotePath': remote, 'contentBase64': __import__('base64').b64encode(buffer.getvalue()).decode(), 'size': buffer.tell()}
            if operation == 'put':
                content = cfg.get('content', ctx['last']); raw = content if isinstance(content, bytes) else (content.encode(cfg.get('encoding', 'utf-8')) if isinstance(content, str) else json.dumps(content).encode()); client.putfo(io.BytesIO(raw), remote); return {'remotePath': remote, 'written': True, 'size': len(raw)}
            if operation == 'delete': client.remove(remote); return {'remotePath': remote, 'deleted': True}
            if operation == 'dir': return {'directory': client.getcwd() or '.', 'entries': [{'name': item.filename, 'size': item.st_size, 'modified': item.st_mtime} for item in client.listdir_attr(remote or '.')]}
            if operation == 'change_dir': client.chdir(remote); return {'directory': client.getcwd()}
            raise RuntimeError(f'Unsupported SFTP operation {operation}')
        finally: client.close(); ssh.close()

    @staticmethod
    def schema_content(schema_id, ctx):
        project = ctx.get('project')
        schema = next((item for item in getattr(project, 'schemas', []) if item.id == schema_id or item.name == schema_id), None) if project and schema_id else None
        return schema.content if schema else ''

    def configured_schema(self, cfg, ctx):
        return cfg.get('schemaText') or self.schema_content(cfg.get('schemaId'), ctx)

    @staticmethod
    def schema_root_name(schema):
        if not schema: return ''
        try:
            document = json.loads(schema)
            return document.get('title') or next(iter((document.get('properties') or {}).keys()), '')
        except (ValueError, TypeError):
            pass
        from xml.etree import ElementTree
        try: root = ElementTree.fromstring(schema)
        except (ElementTree.ParseError, TypeError): return ''
        return next((item.attrib.get('name') for item in list(root) if item.tag.rsplit('}', 1)[-1] == 'element' and item.attrib.get('name')), '')

    @staticmethod
    def schema_leaf_fields(schema):
        if not schema: return [], []
        try:
            document = json.loads(schema); fields = []
            def walk(node):
                properties = node.get('properties', {}) if isinstance(node, dict) else {}
                for name, child in properties.items():
                    if isinstance(child, dict) and child.get('properties'): walk(child)
                    else: fields.append((name, child.get('type', 'string') if isinstance(child, dict) else 'string'))
            walk(document)
            return [item[0] for item in fields], [item[1] for item in fields]
        except (ValueError, TypeError):
            pass
        from xml.etree import ElementTree
        try: document = ElementTree.fromstring(schema)
        except (ElementTree.ParseError, TypeError): return [], []
        fields = []
        def local(element): return element.tag.rsplit('}', 1)[-1]
        def child_elements(element):
            result = []
            for child in list(element):
                if local(child) == 'element': result.append(child)
                elif local(child) in ('complexType', 'sequence', 'all', 'choice', 'group', 'extension'): result.extend(child_elements(child))
            return result
        def walk(element):
            children = child_elements(element)
            if children:
                for child in children: walk(child)
            elif local(element) == 'element' and element.attrib.get('name'):
                fields.append((element.attrib['name'], element.attrib.get('type', 'string').split(':')[-1]))
            for child in element.iter():
                if local(child) == 'attribute' and child.attrib.get('name'): fields.append((child.attrib['name'], child.attrib.get('type', 'string').split(':')[-1]))
        roots = [item for item in list(document) if local(item) == 'element']
        for root in roots: walk(root)
        unique = list(dict.fromkeys(fields))
        return [item[0] for item in unique], [item[1] for item in unique]

    def parse_xml_activity(self, cfg, ctx):
        style = str(cfg.get('inputStyle', 'Text')).lower()
        source = cfg.get('xmlString', cfg.get('source', ctx['last']))
        if style in ('binary','dynamic') and cfg.get('xmlBinary') not in (None, ''): source = cfg.get('xmlBinary')
        if isinstance(source, dict): source = source.get('bytes') or source.get('binaryContent') or source.get('xmlString') or source.get('content') or source.get('body') or ''
        if isinstance(source, str) and style == 'binary':
            try: source = base64.b64decode(source)
            except ValueError: source = source.encode(cfg.get('forceEncoding') or cfg.get('encoding') or 'utf-8')
        if isinstance(source, bytes): source = source.decode(cfg.get('forceEncoding') or cfg.get('encoding') or 'utf-8-sig')
        parsed = self.parse_xml(source)
        if self.as_bool(cfg.get('validateOutput', False)):
            schema = self.configured_schema(cfg, ctx)
            if not schema: raise FabricFault('Validate Output requires a project or inline XSD in Output Editor', fault_type='ValidationException')
            self.validate_xml_root(parsed['root'], schema)
        return parsed

    def render_xml_activity(self, cfg, ctx):
        schema = self.configured_schema(cfg, ctx)
        root_name = self.schema_root_name(schema) or cfg.get('rootElement') or 'root'
        source = cfg.get(root_name, cfg.get('value', cfg.get('source', ctx['last'])))
        if isinstance(source, dict) and root_name in source and not ('root' in source and 'value' in source): source = source[root_name]
        if self.as_bool(cfg.get('validateInput', False)):
            if not schema: raise FabricFault('Validate Input requires a project or inline XSD in Input Editor', fault_type='ValidationException')
            actual_root = source.get('root') if isinstance(source, dict) and 'root' in source else root_name
            self.validate_xml_root(actual_root, schema)
        encoding = cfg.get('encoding') or 'UTF-8'
        content = self.render_xml(source, root_name, 'unicode', pretty=self.as_bool(cfg.get('prettyPrint', False)))
        if not self.as_bool(cfg.get('suppressXmlDeclaration', False)): content = f'<?xml version="1.0" encoding="{encoding}"?>\n{content}'
        if str(cfg.get('outputStyle', 'Text')).lower() == 'binary':
            raw = content.encode(encoding); return {'binaryContent': base64.b64encode(raw).decode(), 'bytes': base64.b64encode(raw).decode(), 'byteCount': len(raw), 'encoding': encoding, 'content': None}
        return {'content': content, 'xmlString': content, 'encoding': encoding}

    @staticmethod
    def validate_xml_root(root_name, schema):
        from xml.etree import ElementTree
        try: schema_root = ElementTree.fromstring(schema)
        except ElementTree.ParseError as exc: raise FabricFault(f'Selected XSD is invalid: {exc}', fault_type='ValidationException') from exc
        expected = next((item.attrib.get('name') for item in list(schema_root) if item.tag.rsplit('}',1)[-1] == 'element' and item.attrib.get('name')), None)
        actual = str(root_name or '').rsplit('}',1)[-1]
        if expected and actual != expected: raise FabricFault(f'XML root {actual!r} does not match XSD root {expected!r}', fault_type='ValidationException', details={'expectedRoot':expected,'actualRoot':actual})

    def parse_xml(self, source):
        from xml.etree import ElementTree
        if isinstance(source, dict): source = source.get('content') or source.get('body') or ''
        root = ElementTree.fromstring(source)
        def convert(element):
            children = list(element); result = {'@' + key: value for key, value in element.attrib.items()}
            if not children:
                text = element.text or ''
                if result:
                    if text: result['#text'] = text
                    return result
                return text
            for child in children:
                value = convert(child)
                if child.tag in result: result[child.tag] = result[child.tag] if isinstance(result[child.tag], list) else [result[child.tag]]; result[child.tag].append(value)
                else: result[child.tag] = value
            if element.text and element.text.strip(): result['#text'] = element.text.strip()
            return result
        value = convert(root)
        xml = ElementTree.tostring(root, encoding='unicode')
        return {
            'mediaType': 'application/xml',
            'xml': xml,
            'xmlString': xml,
            'root': root.tag,
            'value': value,
            root.tag.rsplit('}', 1)[-1]: value,
        }

    def render_xml(self, source, root_name='root', encoding='unicode', pretty=False):
        from xml.etree import ElementTree
        if isinstance(source, dict) and 'root' in source and 'value' in source: root_name, source = source['root'], source['value']
        def build(name, value):
            element = ElementTree.Element(name)
            if isinstance(value, dict):
                for key, item in value.items():
                    if key.startswith('@'): element.set(key[1:], str(item))
                    elif key == '#text': element.text = str(item)
                    elif isinstance(item, list):
                        for entry in item: element.append(build(key, entry))
                    else: element.append(build(key, item))
            elif value is not None: element.text = str(value)
            return element
        root = build(root_name, source)
        if pretty: ElementTree.indent(root, space='  ')
        return ElementTree.tostring(root, encoding=encoding).decode() if encoding != 'unicode' else ElementTree.tostring(root, encoding='unicode')

    def json_activity(self, cfg, ctx):
        operation = cfg.get('operation', 'parse'); source = cfg.get('jsonString', cfg.get('value', cfg.get('source', ctx['last'])))
        if isinstance(source, dict) and operation == 'parse': source = source.get('jsonString') or source.get('content') or source.get('body') or source
        if operation == 'parse':
            if not isinstance(source, (str, bytes, bytearray)): value = source
            else:
                policy = str(cfg.get('duplicateKeyPolicy', 'Last wins')).lower()
                def pairs(items):
                    result = {}
                    for key, value in items:
                        if key in result and policy == 'error': raise ValueError(f'Duplicate JSON key: {key}')
                        if key not in result or policy != 'first wins': result[key] = value
                    return result
                value = json.loads(source, object_pairs_hook=pairs)
            self.validate_json_if_requested(value, cfg, ctx, 'validateOutput')
            return value
        schema = self.configured_schema(cfg, ctx); root_name = self.schema_root_name(schema)
        if root_name and root_name in cfg:
            source = {root_name: cfg[root_name]} if str(cfg.get('rootStyle', 'With root')).lower() != 'anonymous' else cfg[root_name]
        self.validate_json_if_requested(source, cfg, ctx, 'validateInput')
        if self.as_bool(cfg.get('omitNulls', False)):
            def without_nulls(value):
                if isinstance(value, dict): return {key: without_nulls(item) for key, item in value.items() if item is not None}
                if isinstance(value, list): return [without_nulls(item) for item in value if item is not None]
                return value
            source = without_nulls(source)
        indent = int(cfg.get('indent', 2)) if self.as_bool(cfg.get('prettyPrint', True)) else None
        content = json.dumps(source, indent=indent, ensure_ascii=self.as_bool(cfg.get('asciiOnly', False)), separators=None if indent is not None else (',', ':'))
        return {'content': content, 'jsonString': content}

    def validate_json_if_requested(self, value, cfg, ctx, flag):
        if not self.as_bool(cfg.get(flag, False)): return
        content = self.configured_schema(cfg, ctx)
        if not content: raise FabricFault(f'{flag} requires a project or inline schema in the schema editor', fault_type='ValidationException')
        try: schema = json.loads(content)
        except ValueError: return
        if schema.get('type') == 'object' and not isinstance(value, dict): raise FabricFault('JSON value must be an object according to the selected schema', fault_type='ValidationException')
        missing = [key for key in schema.get('required', []) if not isinstance(value, dict) or key not in value]
        if missing: raise FabricFault(f'Missing required JSON fields: {", ".join(missing)}', fault_type='ValidationException', details={'missing':missing})

    def flat_data(self, cfg, ctx):
        source, delimiter = cfg.get('text', cfg.get('records', cfg.get('source', ctx['last']))), str(cfg.get('delimiter', ','))
        if cfg.get('operation') == 'parse' and str(cfg.get('inputSource', 'String')).lower() in ('file', 'file path', 'filepath'):
            file_path = cfg.get('filePath') or (source.get('path') if isinstance(source, dict) else source)
            if not file_path:
                raise FabricFault('Parse Data file input requires a file path', fault_type='ParseDataException')
            try:
                source = Path(str(file_path)).expanduser().read_text(encoding=str(cfg.get('fileEncoding') or cfg.get('encoding') or 'utf-8-sig'))
            except (OSError, UnicodeError) as exc:
                raise FabricFault(f'Unable to read Parse Data input file {file_path!s}: {exc}', fault_type='ParseDataException', details={'filePath': str(file_path)}) from exc
        fields = [item.strip() for item in cfg.get('fields', '').split(',') if item.strip()]
        types = [item.strip().lower() for item in str(cfg.get('fieldTypes', '')).split(',') if item.strip()]
        if not fields:
            schema_fields, schema_types = self.schema_leaf_fields(self.configured_schema(cfg, ctx))
            if schema_fields: fields = schema_fields
            if not types and schema_types: types = [item.lower() for item in schema_types]
        line_ending = {'lf':'\n','crlf':'\r\n','cr':'\r','auto':'\n'}.get(str(cfg.get('lineSeparator', cfg.get('lineEnding', 'Auto'))).lower(), str(cfg.get('lineEnding', '\n')))
        def typed(value, index):
            value = value.strip() if self.as_bool(cfg.get('trimValues', True)) else value
            kind = types[index] if index < len(types) else 'string'
            if value == '': return None if kind not in ('string','text') else ''
            if kind in ('integer','int','long'): return int(value)
            if kind in ('decimal','number','float','double'): return float(value)
            if kind in ('boolean','bool'): return value.lower() in ('true','1','yes','y')
            return value
        def parsed_xml_result(records, names=None):
            from xml.etree import ElementTree
            def xml_name(value, fallback):
                normalized = re.sub(r'[^A-Za-z0-9_.-]+', '_', str(value or fallback))
                return normalized if re.match(r'[A-Za-z_]', normalized) else f'_{normalized}'
            schema = self.configured_schema(cfg, ctx)
            root_name = xml_name(cfg.get('rootElement') or self.schema_root_name(schema), 'records')
            schema_record_name = None
            if schema:
                try:
                    schema_document = ElementTree.fromstring(schema)
                    schema_elements = [item for item in schema_document.iter() if item.tag.rsplit('}', 1)[-1] == 'element' and item.attrib.get('name')]
                    schema_record_name = schema_elements[1].attrib['name'] if len(schema_elements) > 1 else None
                except (ElementTree.ParseError, TypeError):
                    pass
            record_name = xml_name(cfg.get('recordElement') or schema_record_name, 'record')
            root = ElementTree.Element(root_name)
            for record in records:
                item = ElementTree.SubElement(root, record_name)
                values = record if isinstance(record, dict) else {'value': record}
                for key, value in values.items():
                    child = ElementTree.SubElement(item, xml_name(key, 'field'))
                    if value is not None: child.text = str(value).lower() if isinstance(value, bool) else str(value)
            xml = ElementTree.tostring(root, encoding='unicode')
            return {
                'mediaType': 'application/xml', 'xml': xml, 'xmlString': xml,
                'root': root_name, 'value': {record_name: records},
                'records': records, 'recordCount': len(records),
                **({'fields': names} if names else {}),
            }
        if str(cfg.get('format', 'delimited')).lower() == 'fixed':
            widths = [int(item.strip()) for item in cfg.get('widths', '').split(',') if item.strip()]
            if not fields or len(fields) != len(widths): raise RuntimeError('Fixed-width data requires matching field names and widths')
            if cfg.get('operation') == 'parse':
                records = []
                for line in str(source).splitlines():
                    if not line and self.as_bool(cfg.get('skipBlankLines', True)): continue
                    if self.as_bool(cfg.get('strictColumns', False)) and len(line) != sum(widths): raise FabricFault(f'Fixed-width line length {len(line)} does not match configured width {sum(widths)}', fault_type='ParseDataException')
                    offset, record = 0, {}
                    for index, (name, width) in enumerate(zip(fields, widths)): record[name], offset = typed(line[offset:offset + width].rstrip(str(cfg.get('fillCharacter', ' ')) or ' '), index), offset + width
                    records.append(record)
                return parsed_xml_result(records, fields)
            records = source.get('records', source) if isinstance(source, dict) else source
            if not isinstance(records, list): records = [records]
            fill = str(cfg.get('fillCharacter', ' ') or ' ')[0]
            def fixed_field(record, name, width, index):
                value = '' if record.get(name) is None else str(record.get(name)); value = value[:width]
                return value.rjust(width, fill) if index < len(types) and types[index] in ('integer','int','long','decimal','number','float','double') else value.ljust(width, fill)
            content = line_ending.join(''.join(fixed_field(record, name, width, index) for index, (name, width) in enumerate(zip(fields, widths))) for record in records)
            if self.as_bool(cfg.get('includeFinalLineSeparator', False)): content += line_ending
            return {'content': content, 'recordCount': len(records)}
        if cfg.get('operation') == 'parse':
            text = str(source.get('content', source) if isinstance(source, dict) else source)
            lines = [line for line in text.splitlines() if line or not self.as_bool(cfg.get('skipBlankLines', True))]
            if len(delimiter) == 1:
                rows = list(csv.reader(lines, delimiter=delimiter))
            else:
                pattern = f'[{re.escape(delimiter)}]' if str(cfg.get('separatorRule', 'single')).lower() == 'any-character' else re.escape(delimiter)
                rows = [re.split(pattern, line) for line in lines]
            if self.as_bool(cfg.get('header', True)) and rows: names, rows = rows[0], rows[1:]
            else: names = fields
            if not names: raise FabricFault('Parse Data requires header fields or configured field names', fault_type='ParseDataException')
            records = []
            for row in rows:
                if self.as_bool(cfg.get('strictColumns', False)) and len(row) != len(names): raise FabricFault(f'Column count {len(row)} does not match schema field count {len(names)}', fault_type='ParseDataException')
                records.append({name: typed(row[index] if index < len(row) else '', index) for index, name in enumerate(names)})
            return parsed_xml_result(records, names)
        records = source.get('records', source) if isinstance(source, dict) else source
        if not isinstance(records, list): records = [records]
        names = fields or list(records[0].keys() if records else [])
        if len(delimiter) == 1:
            output = io.StringIO(); writer = csv.DictWriter(output, fieldnames=names, delimiter=delimiter, lineterminator=line_ending)
            if self.as_bool(cfg.get('header', True)): writer.writeheader()
            writer.writerows(records); content = output.getvalue()
            if not self.as_bool(cfg.get('includeFinalLineSeparator', True)): content = content.removesuffix(line_ending)
        else:
            rows = ([delimiter.join(names)] if self.as_bool(cfg.get('header', True)) else []) + [delimiter.join('' if record.get(name) is None else str(record.get(name)) for name in names) for record in records]
            content = line_ending.join(rows) + (line_ending if self.as_bool(cfg.get('includeFinalLineSeparator', False)) else '')
        return {'content': content, 'recordCount': len(records), 'fields':names}

    async def java_worker(self, cfg, payload):
        command = cfg.get('command') or os.getenv('JAVA_WORKER_COMMAND')
        if command:
            proc = await asyncio.create_subprocess_shell(command, stdin=asyncio.subprocess.PIPE, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE)
            out, err = await proc.communicate(json.dumps({'className': cfg.get('className'), 'method': cfg.get('method'), 'payload': payload}).encode())
            if proc.returncode: raise RuntimeError(err.decode() or 'Java worker failed')
            return json.loads(out.decode())
        class_name, method = str(cfg.get('className') or '').strip(), str(cfg.get('method') or '').strip()
        if not class_name or not method: raise RuntimeError('Java Invoke requires a class name and method')
        artifact = Path(str(cfg.get('artifactPath') or '')).expanduser()
        source = str(cfg.get('sourceCode') or '')
        parameters = cfg.get('parameters')
        if not isinstance(parameters, list): parameters = [cfg.get('payload', payload)]
        helper = '''import java.lang.reflect.*; public class FabricInvoker { public static void main(String[] a) throws Exception { Class<?> c=Class.forName(a[0]); Method found=null; for(Method m:c.getMethods()) if(m.getName().equals(a[1])&&m.getParameterCount()==a.length-2){found=m;break;} if(found==null) throw new NoSuchMethodException(a[0]+"."+a[1]); Object[] v=new Object[found.getParameterCount()]; Class<?>[] t=found.getParameterTypes(); for(int i=0;i<v.length;i++){String s=a[i+2]; v[i]=t[i]==String.class?s:t[i]==int.class||t[i]==Integer.class?Integer.valueOf(s):t[i]==long.class||t[i]==Long.class?Long.valueOf(s):t[i]==double.class||t[i]==Double.class?Double.valueOf(s):t[i]==boolean.class||t[i]==Boolean.class?Boolean.valueOf(s):s;} Object target=Modifier.isStatic(found.getModifiers())?null:c.getDeclaredConstructor().newInstance(); Object out=found.invoke(target,v); if(out!=null) System.out.print(out); }}'''
        with tempfile.TemporaryDirectory(prefix='fabric-java-') as folder:
            root = Path(folder); (root / 'FabricInvoker.java').write_text(helper, encoding='utf-8')
            classpath = str(root)
            compile_inputs = [str(root / 'FabricInvoker.java')]
            if source:
                java_file = root / f'{class_name.rsplit(".", 1)[-1]}.java'; java_file.write_text(source, encoding='utf-8'); compile_inputs.append(str(java_file))
            elif artifact.exists() and artifact.suffix.lower() == '.java': compile_inputs.append(str(artifact)); classpath += os.pathsep + str(artifact.parent)
            elif artifact.exists() and artifact.suffix.lower() == '.jar': classpath += os.pathsep + str(artifact)
            elif artifact.exists() and artifact.suffix.lower() == '.class': classpath += os.pathsep + str(artifact.parent)
            elif not artifact.exists(): raise RuntimeError('Java Invoke requires an existing JAR/class/source artifact or inline source')
            compiler = await asyncio.create_subprocess_exec('javac', '-cp', classpath, '-d', str(root), *compile_inputs, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE)
            _, compile_error = await compiler.communicate()
            if compiler.returncode: raise RuntimeError(f'Java compilation failed: {compile_error.decode().strip()}')
            args = [json.dumps(value, separators=(',', ':')) if isinstance(value, (dict, list)) else str(value) for value in parameters]
            process = await asyncio.create_subprocess_exec('java', '-cp', classpath, 'FabricInvoker', class_name, method, *args, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE)
            try: out, err = await asyncio.wait_for(process.communicate(), timeout=float(cfg.get('timeout') or 60))
            except asyncio.TimeoutError: process.kill(); raise RuntimeError('Java method invocation timed out')
            if process.returncode: raise RuntimeError(err.decode().strip() or 'Java method invocation failed')
            value = out.decode().strip()
            try: value = json.loads(value)
            except ValueError: pass
            return {'methodReturnValue': value, 'className': class_name, 'method': method}

    async def python_worker(self, cfg, payload):
        def invoke():
            function_name = str(cfg.get('function') or '').strip()
            if not function_name: raise RuntimeError('Python Invoke requires a function name')
            source, artifact = str(cfg.get('sourceCode') or ''), Path(str(cfg.get('artifactPath') or '')).expanduser()
            module_name = str(cfg.get('moduleName') or artifact.stem or 'fabric_inline')
            if source:
                namespace = {'__name__': module_name}; exec(compile(source, f'<{module_name}>', 'exec'), namespace); function = namespace.get(function_name)
            elif artifact.suffix.lower() == '.py' and artifact.exists():
                spec = importlib.util.spec_from_file_location(module_name, artifact)
                if not spec or not spec.loader: raise RuntimeError(f'Cannot load Python module {artifact}')
                module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module); function = getattr(module, function_name, None)
            elif artifact.exists():
                sys.path.insert(0, str(artifact));
                try: module = __import__(module_name, fromlist=[function_name]); function = getattr(module, function_name, None)
                finally: sys.path.pop(0)
            else: raise RuntimeError('Python Invoke requires an existing .py/package artifact or inline source')
            if not callable(function): raise RuntimeError(f'Python function {module_name}.{function_name} was not found')
            parameters = cfg.get('parameters')
            if isinstance(parameters, list): result = function(*parameters)
            elif isinstance(parameters, dict): result = function(**parameters)
            else: result = function(cfg.get('payload', payload))
            return {'result': result, 'module': module_name, 'function': function_name}
        return await asyncio.wait_for(asyncio.to_thread(invoke), timeout=float(cfg.get('timeout') or 60))
